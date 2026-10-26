import sys
sys.stdout.reconfigure(encoding='utf-8')
import docx

doc = docx.Document("Ejercicios U Mann-Whitney.docx")
body = doc._body._element

print(f"Total paragraphs: {len(doc.paragraphs)}")
print(f"Total tables: {len(doc.tables)}")

tbl_idx = 0
p_idx = 0
total_imgs = 0

for i, child in enumerate(body):
    tag = child.tag.split("}")[-1]
    if tag == "tbl":
        t = doc.tables[tbl_idx]
        tbl_idx += 1
        cells_txt = [c.text.strip().replace("\n", " ") for c in t.rows[0].cells[:4]]
        print(f"[{i:02d}] TABLE {tbl_idx-1}: {len(t.rows)}x{len(t.columns)} cols -> {cells_txt}")
    elif tag == "p":
        p = doc.paragraphs[p_idx]
        p_idx += 1
        imgs = len(p._p.xpath('.//@r:embed'))
        total_imgs += imgs
        t = p.text.strip()
        if t or imgs:
            print(f"[{i:02d}] P{p_idx-1:02d} [imgs={imgs}]: {t[:85]}")

print(f"\nTotal images found: {total_imgs}")
