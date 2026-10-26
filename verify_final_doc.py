import sys
sys.stdout.reconfigure(encoding='utf-8')
import docx

doc = docx.Document("Ejercicios Pruebas de chi cuadrado.docx")
print(f"Total paragraphs: {len(doc.paragraphs)}")
print(f"Total tables: {len(doc.tables)}")

total_images = 0
for i, p in enumerate(doc.paragraphs):
    imgs = len(p._p.xpath('.//@r:embed'))
    total_images += imgs
    t = p.text.strip()
    if t or imgs:
        print(f"P{i:02d} [imgs={imgs}]: {t[:85]}")

print(f"\nTotal images found in document: {total_images}")
