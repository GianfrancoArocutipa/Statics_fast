# -*- coding: utf-8 -*-
import os
import sys
import docx
import zipfile
import pandas as pd

def inspect_docx(filename):
    print(f"=================== {filename} ===================")
    doc = docx.Document(filename)
    for i, p in enumerate(doc.paragraphs):
        text = p.text.strip()
        num_drawings = len([r for r in p.runs if 'drawing' in r._r.xml])
        if text or num_drawings > 0:
            print(f"[{i}] {text} (drawings: {num_drawings})")
            
    print(f"\nTables count: {len(doc.tables)}")
    for t_idx, table in enumerate(doc.tables):
        print(f"--- Table {t_idx} ({len(table.rows)} rows x {len(table.columns)} cols) ---")
        for r_idx, row in enumerate(table.rows[:5]):
            cells = [c.text.strip().replace('\n', ' ') for c in row.cells]
            print(f"  R{r_idx}: {cells}")

if __name__ == '__main__':
    with open('inspect_model.txt', 'w', encoding='utf-8') as out:
        sys.stdout = out
        inspect_docx('PEREZ _AROCUTIPA_LLICA - APLICACIÓN DE ESTADÍSTICA INFERENCIAL CON DATOS REALES.docx')
        inspect_docx('trabajo.docx')
        inspect_docx('Ejerc_prueba ANOVA.docx')
