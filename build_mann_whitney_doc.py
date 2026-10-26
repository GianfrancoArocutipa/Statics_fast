import os
import shutil
import docx
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

original_file = "Ejercicios U Mann-Whitney.docx"
backup_file = "Ejercicios U Mann-Whitney_Original.docx"
output_file = "Ejercicios U Mann-Whitney_Resuelto.docx"

if not os.path.exists(backup_file):
    shutil.copy(original_file, backup_file)
    print(f"Created backup: {backup_file}")

doc = docx.Document(original_file)

def add_p_before(target_p, text, is_empty=False):
    p = target_p.insert_paragraph_before(text if not is_empty else "", style="List Paragraph")
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    for r in p.runs:
        r.font.name = "Calibri"
        r.font.size = Pt(11)
    return p

def add_img_before(target_p, img_path, width):
    p = target_p.insert_paragraph_before("", style="List Paragraph")
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run()
    r.add_picture(img_path, width=width)
    return p

def append_p(doc, text, is_empty=False):
    p = doc.add_paragraph(text if not is_empty else "", style="List Paragraph")
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    for r in p.runs:
        r.font.name = "Calibri"
        r.font.size = Pt(11)
    return p

def append_img(doc, img_path, width):
    p = doc.add_paragraph("", style="List Paragraph")
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run()
    r.add_picture(img_path, width=width)
    return p

# 1. Exercise 1: Insert before P6 ("Se están comparando dos sistemas de almacenamiento")
target_ex2 = None
for p in doc.paragraphs:
    if "Se están comparando dos sistemas de almacenamiento" in p.text:
        target_ex2 = p
        break

if not target_ex2:
    raise Exception("Could not find target paragraph for Exercise 1 insertion!")

print("Inserting Exercise 1 resolution...")
add_p_before(target_ex2, "Planteamiento de Hipótesis")
add_p_before(target_ex2, "H0: No existe diferencia significativa en el tiempo de ejecución del algoritmo de búsqueda entre los lenguajes de programación Java y Python (las medianas/distribuciones son iguales).")
add_p_before(target_ex2, "", is_empty=True)
add_p_before(target_ex2, "H1: Existe una diferencia significativa en el tiempo de ejecución del algoritmo de búsqueda entre los lenguajes de programación Java y Python (las medianas/distribuciones son diferentes).")
add_p_before(target_ex2, "", is_empty=True)
add_p_before(target_ex2, "Determinar el nivel de Significancia:5% (0.05)")
add_p_before(target_ex2, "Tabla de rangos")
add_img_before(target_ex2, "generated_spss_mann_images/spss_mann_rangos_ex1.png", width=Inches(3.8))
add_p_before(target_ex2, "", is_empty=True)
add_p_before(target_ex2, "Seleccionar el estadístico de prueba: Prueba no paramétrica U de Mann-Whitney")
add_img_before(target_ex2, "generated_spss_mann_images/spss_mann_test_ex1.png", width=Inches(3.8))
add_p_before(target_ex2, "Contrastación de Hipótesis")
add_p_before(target_ex2, "Dado que p-valor (0.008) es menor que el nivel de significancia (0.05), es posible que se rechace la H0")
add_p_before(target_ex2, "Conclusión")
add_p_before(target_ex2, "A un margen del 5% se concluye que, Existe una diferencia significativa en el tiempo de ejecución del algoritmo de búsqueda entre los lenguajes de programación Java y Python, siendo Python significativamente más rápido al presentar menores tiempos de ejecución (rango promedio de 3.30 frente a 9.31 de Java).")
add_p_before(target_ex2, "", is_empty=True)

# 2. Exercise 2: Insert before P9 ("Supongamos que un equipo de ingeniería de sistemas está comparando la velocidad de procesamiento de dos servidores diferentes, A y B")
target_ex3 = None
for p in doc.paragraphs:
    if "Supongamos que un equipo de ingeniería de sistemas está comparando la velocidad de procesamiento de dos servidores diferentes, A y B" in p.text:
        target_ex3 = p
        break

if not target_ex3:
    raise Exception("Could not find target paragraph for Exercise 2 insertion!")

print("Inserting Exercise 2 resolution...")
add_p_before(target_ex3, "Planteamiento de Hipótesis")
add_p_before(target_ex3, "H0: El tiempo de acceso a los archivos del Sistema A no es menor que el tiempo de acceso a los archivos del Sistema B en términos de milisegundos.")
add_p_before(target_ex3, "", is_empty=True)
add_p_before(target_ex3, "H1: El tiempo de acceso a los archivos del Sistema A es menor que el tiempo de acceso a los archivos del Sistema B en términos de milisegundos.")
add_p_before(target_ex3, "", is_empty=True)
add_p_before(target_ex3, "Determinar el nivel de Significancia:3% (0.03)")
add_p_before(target_ex3, "Tabla de rangos")
add_img_before(target_ex3, "generated_spss_mann_images/spss_mann_rangos_ex2.png", width=Inches(3.8))
add_p_before(target_ex3, "", is_empty=True)
add_p_before(target_ex3, "Seleccionar el estadístico de prueba: Prueba no paramétrica U de Mann-Whitney")
add_img_before(target_ex3, "generated_spss_mann_images/spss_mann_test_ex2.png", width=Inches(3.8))
add_p_before(target_ex3, "Contrastación de Hipótesis")
add_p_before(target_ex3, "Dado que p-valor (0.001) es menor que el nivel de significancia (0.03), es posible que se rechace la H0")
add_p_before(target_ex3, "Conclusión")
add_p_before(target_ex3, "A un margen del 3% se concluye que, El tiempo de acceso a los archivos del Sistema A es significativamente menor que el tiempo de acceso a los archivos del Sistema B en términos de milisegundos (ofreciendo una velocidad de acceso superior).")
add_p_before(target_ex3, "", is_empty=True)

# 3. Exercise 3: Insert before P12 ("Supongamos que una empresa está investigando el rendimiento de dos algoritmos de compresión")
target_ex4 = None
for p in doc.paragraphs:
    if "Supongamos que una empresa está investigando el rendimiento de dos algoritmos de compresión" in p.text:
        target_ex4 = p
        break

if not target_ex4:
    raise Exception("Could not find target paragraph for Exercise 3 insertion!")

print("Inserting Exercise 3 resolution...")
add_p_before(target_ex4, "Planteamiento de Hipótesis")
add_p_before(target_ex4, "H0: No existe diferencia significativa en el tiempo de procesamiento entre el Servidor A y el Servidor B para manejar solicitudes de base de datos.")
add_p_before(target_ex4, "", is_empty=True)
add_p_before(target_ex4, "H1: Existe una diferencia significativa en el tiempo de procesamiento entre el Servidor A y el Servidor B para manejar solicitudes de base de datos.")
add_p_before(target_ex4, "", is_empty=True)
add_p_before(target_ex4, "Determinar el nivel de Significancia:3% (0.03)")
add_p_before(target_ex4, "Tabla de rangos")
add_img_before(target_ex4, "generated_spss_mann_images/spss_mann_rangos_ex3.png", width=Inches(3.8))
add_p_before(target_ex4, "", is_empty=True)
add_p_before(target_ex4, "Seleccionar el estadístico de prueba: Prueba no paramétrica U de Mann-Whitney")
add_img_before(target_ex4, "generated_spss_mann_images/spss_mann_test_ex3.png", width=Inches(3.8))
add_p_before(target_ex4, "Contrastación de Hipótesis")
add_p_before(target_ex4, "Dado que p-valor (0.000) es menor que el nivel de significancia (0.03), es posible que se rechace la H0")
add_p_before(target_ex4, "Conclusión")
add_p_before(target_ex4, "A un margen del 3% se concluye que, Existe una diferencia significativa en el tiempo de procesamiento entre los dos servidores para manejar solicitudes de base de datos, siendo el Servidor A significativamente más rápido (rango promedio de 11.00 frente a 28.94 del Servidor B).")
add_p_before(target_ex4, "", is_empty=True)

# 4. Exercise 4: Insert after Table 5 / append at the end
print("Inserting Exercise 4 resolution...")
append_p(doc, "Planteamiento de Hipótesis")
append_p(doc, "H0: No existe diferencia significativa en la tasa de compresión entre el Algoritmo A y el Algoritmo B.")
append_p(doc, "", is_empty=True)
append_p(doc, "H1: Existe una diferencia significativa en la tasa de compresión entre el Algoritmo A y el Algoritmo B.")
append_p(doc, "", is_empty=True)
append_p(doc, "Determinar el nivel de Significancia:5% (0.05)")
append_p(doc, "Tabla de rangos")
append_img(doc, "generated_spss_mann_images/spss_mann_rangos_ex4.png", width=Inches(3.8))
append_p(doc, "", is_empty=True)
append_p(doc, "Seleccionar el estadístico de prueba: Prueba no paramétrica U de Mann-Whitney")
append_img(doc, "generated_spss_mann_images/spss_mann_test_ex4.png", width=Inches(3.8))
append_p(doc, "Contrastación de Hipótesis")
append_p(doc, "Dado que p-valor (0.000) es menor que el nivel de significancia (0.05), es posible que se rechace la H0")
append_p(doc, "Conclusión")
append_p(doc, "A un margen del 5% se concluye que, Existe una diferencia significativa en la tasa de compresión entre los dos algoritmos de compresión de datos, logrando el Algoritmo B una tasa de compresión significativamente superior a la del Algoritmo A (rango promedio de 26.73 frente a 11.45 del Algoritmo A).")

# Save to output file and update original file
doc.save(output_file)
print(f"Saved completed report to {output_file}")
doc.save(original_file)
print(f"Updated original file {original_file}")
