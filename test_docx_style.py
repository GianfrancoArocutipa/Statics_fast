import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

doc = docx.Document()
p = doc.add_paragraph('Test Paragraph')
table = doc.add_table(rows=2, cols=2)
shading_elm = parse_xml(r'<w:shd {} w:fill="1F497D"/>'.format(nsdecls('w')))
table.cell(0, 0)._tc.get_or_add_tcPr().append(shading_elm)
doc.save('test_out.docx')
print('Table formatting test passed!')
