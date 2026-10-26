# -*- coding: utf-8 -*-
"""
Script to build the streamlined, highly rigorous Word document following:
1. Exact narrative style of the engineer in 'Ejerc_prueba ANOVA.docx':
   - Paso 1: Planteamiento y Pregunta de investigación
   - Paso 2: Formulación de Hipótesis (H0 y H1) y nivel alfa = 0.05
   - Paso 3: Estadísticos descriptivos de cada grupo (Media, Desv. Estándar, n)
   - Paso 4: Verificación de Supuestos (Shapiro-Wilk para normalidad y Levene para homocedasticidad)
   - Paso 5: Aplicación de la prueba inferencial (estadístico t o F, gl, p-valor)
   - Paso 6: Regla y Decisión Estadística (p >= alfa -> H0; p < alfa -> H1)
   - Paso 7: Pruebas Post Hoc (Tukey o Games-Howell según Levene)
   - Paso 8: Conclusión e interpretación en el contexto de Ingeniería de Sistemas (con tamaño del efecto).
2. Page count strictly <= 25 pages (aiming for 18-22 pages).
3. Incorporates the critical images from 'trabajo.docx', 'Reporte_Estadistico (1).docx', and the model document.
"""

import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def generate_streamlined_doc():
    doc = docx.Document()

    # 1. Page Margins & Setup (Standard 1 inch / 2.54 cm)
    for sec in doc.sections:
        sec.top_margin = Inches(1.0)
        sec.bottom_margin = Inches(1.0)
        sec.left_margin = Inches(1.0)
        sec.right_margin = Inches(1.0)
        sec.page_width = Inches(8.5)
        sec.page_height = Inches(11.0)

    # 2. Typography & Styles
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Times New Roman'
    style_normal.font.size = Pt(11) # 11pt ensures tight, professional page budget <= 25 pages
    style_normal.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    style_normal.paragraph_format.line_spacing = 1.15
    style_normal.paragraph_format.space_after = Pt(3)
    style_normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    # Institutional Palette
    NAVY = RGBColor(0x1B, 0x36, 0x5D)
    DARK_BLUE = RGBColor(0x2C, 0x3E, 0x50)
    SLATE = RGBColor(0x34, 0x49, 0x5E)
    BODY = RGBColor(0x22, 0x22, 0x22)
    GRAY = RGBColor(0x55, 0x55, 0x55)

    def add_p(text="", bold=False, italic=False, size=11, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_before=0, space_after=3, color=BODY):
        p = doc.add_paragraph()
        p.alignment = align
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.15
        if text:
            run = p.add_run(text)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.italic = italic
            run.font.color.rgb = color
        return p

    def add_h1(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.color.rgb = NAVY
        return p

    def add_h2(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(7)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(11.5)
        run.font.bold = True
        run.font.color.rgb = DARK_BLUE
        return p

    def add_h3(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(5)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(11)
        run.font.bold = True
        run.font.italic = True
        run.font.color.rgb = SLATE
        return p

    def add_img(img_path, caption, width_in=4.2):
        if os.path.exists(img_path):
            p_img = doc.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.paragraph_format.space_before = Pt(5)
            p_img.paragraph_format.space_after = Pt(2)
            p_img.paragraph_format.keep_with_next = True
            run_img = p_img.add_run()
            run_img.add_picture(img_path, width=Inches(width_in))

            p_cap = doc.add_paragraph()
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cap.paragraph_format.space_before = Pt(1)
            p_cap.paragraph_format.space_after = Pt(6)
            run_cap = p_cap.add_run(caption)
            run_cap.font.name = 'Times New Roman'
            run_cap.font.size = Pt(9.5)
            run_cap.font.bold = True
            run_cap.font.italic = True
            run_cap.font.color.rgb = GRAY
        else:
            print(f"Warning: Image {img_path} not found.")

    def create_custom_table(data, col_widths=None, header_bg="1B365D", alt_bg="F8F9F9"):
        table = doc.add_table(rows=len(data), cols=len(data[0]))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER

        for r_idx, row in enumerate(table.rows):
            trPr = row._tr.get_or_add_trPr()
            trPr.append(parse_xml(r'<w:cantSplit {}/>'.format(nsdecls('w'))))
            if r_idx == 0:
                trPr.append(parse_xml(r'<w:tblHeader {}/>'.format(nsdecls('w'))))

            for c_idx, cell in enumerate(row.cells):
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                cell.text = str(data[r_idx][c_idx])
                if col_widths and c_idx < len(col_widths):
                    cell.width = col_widths[c_idx]

                tcPr = cell._tc.get_or_add_tcPr()
                if r_idx == 0:
                    shd = parse_xml(r'<w:shd {} w:fill="{}"/>'.format(nsdecls('w'), header_bg))
                    tcPr.append(shd)
                    for cp in cell.paragraphs:
                        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        for run in cp.runs:
                            run.font.name = 'Times New Roman'
                            run.font.bold = True
                            run.font.size = Pt(9.5)
                            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                else:
                    if r_idx % 2 == 1:
                        shd = parse_xml(r'<w:shd {} w:fill="FFFFFF"/>'.format(nsdecls('w')))
                    else:
                        shd = parse_xml(r'<w:shd {} w:fill="{}"/>'.format(nsdecls('w'), alt_bg))
                    tcPr.append(shd)
                    for cp in cell.paragraphs:
                        if c_idx == 0:
                            cp.alignment = WD_ALIGN_PARAGRAPH.LEFT
                        else:
                            cp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                        for run in cp.runs:
                            run.font.name = 'Times New Roman'
                            run.font.size = Pt(9)
                            run.font.color.rgb = BODY

                borders = parse_xml(
                    r'<w:tcBorders {} >'
                    r'  <w:top w:val="single" w:sz="4" w:space="0" w:color="BDC3C7"/>'
                    r'  <w:left w:val="none"/>'
                    r'  <w:bottom w:val="single" w:sz="4" w:space="0" w:color="BDC3C7"/>'
                    r'  <w:right w:val="none"/>'
                    r'</w:tcBorders>'.format(nsdecls('w'))
                )
                tcPr.append(borders)

        p_after = doc.add_paragraph()
        p_after.paragraph_format.space_before = Pt(0)
        p_after.paragraph_format.space_after = Pt(4)
        return table

    # =========================================================================
    # CARÁTULA
    # =========================================================================
    add_p("UNIVERSIDAD PRIVADA DE TACNA", bold=True, size=15, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=10, space_after=3, color=NAVY)
    add_p("FACULTAD DE INGENIERÍA", bold=True, size=13, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2, color=NAVY)
    add_p("ESCUELA DE INGENIERÍA DE SISTEMAS", bold=True, size=12, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12, color=NAVY)

    logo_path = 'extracted_model_media/image8.png'
    if os.path.exists(logo_path):
        p_logo = doc.add_paragraph()
        p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_logo.paragraph_format.space_after = Pt(12)
        run_logo = p_logo.add_run()
        run_logo.add_picture(logo_path, width=Inches(2.2))

    add_p("Trabajo Práctico", bold=True, size=12, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=3, color=DARK_BLUE)
    add_p("“APLICACIÓN DE ESTADÍSTICA INFERENCIAL CON DATOS REALES”", bold=True, size=14, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=16, color=NAVY)

    add_p("Que se presenta para el curso:", italic=True, size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    add_p("ESTADÍSTICA INFERENCIAL", bold=True, size=12, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=14, color=DARK_BLUE)

    add_p("Docente:", bold=True, size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    add_p("LUIS ALFREDO FERNANDEZ VIZCARRA", bold=True, size=11.5, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=16)

    add_p("Integrantes:", bold=True, size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=3)
    add_p("Arocutipa Arocutipa, Gian Franco  (2023076790)", size=11, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    add_p("Perez Peralta, Fabrizio Salvador Elias  (2023077476)", size=11, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    add_p("LLica Mamani, Jimmy Mijair  (2023076789)", size=11, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=20)

    add_p("TACNA – PERÚ", bold=True, size=11, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    add_p("2026", bold=True, size=11, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)

    doc.add_page_break()

    # =========================================================================
    # ÍNDICE GENERAL (Compact)
    # =========================================================================
    add_p("Índice General", bold=True, size=14, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=6, space_after=10, color=NAVY)

    toc_items = [
        ("1. Descripción del dataset", "2"),
        ("2. Planteamiento del análisis y preguntas de investigación", "2"),
        ("3. Análisis descriptivo y estimación puntual e intervalar", "3"),
        ("4. Verificación de supuestos: Normalidad y Homogeneidad de Varianzas", "4"),
        ("5. Procedimiento 1: Prueba de hipótesis para una muestra (stars)", "6"),
        ("6. Procedimiento 2: Comparación de promedios entre dos grupos independientes (ln_forks_1)", "8"),
        ("7. Procedimiento 3: Comparación de promedios en muestras relacionadas (ln_stars_1 vs. ln_forks_1)", "11"),
        ("8. Procedimiento 4: Análisis de Varianza (ANOVA) para tres grupos independientes (ln_open_issues_1)", "13"),
        ("9. Procedimiento 5: Pruebas Post Hoc de comparaciones múltiples (HSD Tukey)", "16"),
        ("10. Conclusiones generales", "18"),
        ("11. Referencias y fuentes de datos", "19")
    ]

    for item, page in toc_items:
        p_toc = doc.add_paragraph()
        p_toc.paragraph_format.space_before = Pt(1)
        p_toc.paragraph_format.space_after = Pt(2)
        p_toc.paragraph_format.line_spacing = 1.15
        r_item = p_toc.add_run(item)
        r_item.font.name = 'Times New Roman'
        r_item.font.size = Pt(10)
        if item.startswith("1.") or item.startswith("5.") or item.startswith("6.") or item.startswith("7.") or item.startswith("8.") or item.startswith("9.") or item.startswith("10.") or item.startswith("11."):
            r_item.font.bold = True
            r_item.font.color.rgb = NAVY
        else:
            r_item.font.color.rgb = BODY

    doc.add_page_break()

    # =========================================================================
    # 1. DESCRIPCIÓN DEL DATASET
    # =========================================================================
    add_h1("1. Descripción del dataset")
    add_p(
        "El presente informe práctico aplica los conceptos y procedimientos de la inferencia estadística sobre el dataset real "
        "titulado “GitHub Repositories – Heterogeneous Data”, obtenido del repositorio público Kaggle Datasets (https://www.kaggle.com/datasets). "
        "El archivo fuente contiene 66 641 registros heterogéneos vinculados al desarrollo de software abierto. Para este estudio se filtró la entidad "
        "RepositoryItem correspondiente a repositorios cuyo lenguaje primario es Python, consolidando una muestra representativa de n = 1567 "
        "proyectos independientes."
    )
    add_p("Variables analizadas:")
    add_p("• stars (Cuantitativa discreta): Número de estrellas recibidas; principal indicador de popularidad e impacto comunitario.")
    add_p("• forks (Cuantitativa discreta): Número de copias derivadas; cuantifica la adopción técnica y reutilización del código.")
    add_p("• open_issues (Cuantitativa discreta): Cantidad de incidencias abiertas; mide la carga de mantenimiento y deuda técnica pendiente.")
    add_p("• ln_stars_1, ln_forks_1, ln_open_issues_1 (Cuantitativas transformadas): Transformaciones logarítmicas ln(X + 1) para estabilizar varianzas.")
    add_p("• owner_type / owner2 (Cualitativa nominal dicotómica): Tipo de propietario: Organization (n = 878) y User (n = 689).")
    add_p("• license_spdx_id / lic3 (Cualitativa nominal politómica): Licencia del software: MIT (lic3 = 1.00; n = 546), Apache-2.0 (lic3 = 2.00; n = 405) y GPL-3.0 (lic3 = 3.00; n = 112), sumando N = 1063 para el ANOVA.")

    # =========================================================================
    # 2. PLANTEAMIENTO DEL ANÁLISIS
    # =========================================================================
    add_h1("2. Planteamiento del análisis y preguntas de investigación")
    add_p(
        "En Ingeniería de Sistemas, el éxito y sostenibilidad de una solución de código abierto dependen de factores organizacionales y jurídicos. "
        "Se busca evaluar si los indicadores de popularidad, reutilización y mantenimiento presentan diferencias significativas o si se ajustan a umbrales estándar. "
        "Para ello, se formulan las siguientes 5 preguntas de investigación guiadas por la narrativa del curso:"
    )
    add_p("• Q1 (Una muestra): ¿El promedio poblacional de popularidad (stars) de los repositorios de Python difiere del estándar de 9000 estrellas?")
    add_p("• Q2 (Dos grupos independientes): ¿Existe diferencia significativa en las bifurcaciones transformadas (ln_forks_1) entre proyectos de Organizaciones y de Usuarios individuales?")
    add_p("• Q3 (Muestras relacionadas): ¿Existe diferencia estadísticamente significativa entre el promedio de estrellas (ln_stars_1) y bifurcaciones (ln_forks_1) evaluadas sobre los mismos repositorios?")
    add_p("• Q4 (ANOVA de un factor): ¿Existen diferencias significativas en las incidencias abiertas transformadas (ln_open_issues_1) según el tipo de licencia (MIT, Apache-2.0 y GPL-3.0)?")
    add_p("• Q5 (Pruebas Post Hoc): ¿Entre qué grupos específicos de licencias se localizan las diferencias significativas en la carga de mantenimiento?")

    # =========================================================================
    # 3. ANÁLISIS DESCRIPTIVO Y ESTIMACIÓN
    # =========================================================================
    add_h1("3. Análisis descriptivo y estimación puntual e intervalar")
    add_p(
        "La Tabla 1 resume los estadísticos descriptivos fundamentales, asimetría, curtosis e intervalos de confianza al 95 % "
        "para la muestra completa (n = 1567). En todas las variables la media supera a la mediana, constatando asimetría positiva."
    )

    add_p("Tabla 1. Estadísticos descriptivos generales e intervalos de confianza al 95 % (n = 1567).", bold=True, size=10, align=WD_ALIGN_PARAGRAPH.CENTER, color=NAVY)
    t1_data = [
        ["Variable", "Media (μ̂)", "Desv. Est. (s)", "Mediana", "Mínimo", "Máximo", "Asimetría", "Curtosis", "IC 95 % para la Media"],
        ["Stars", "8963.88", "4411.51", "8539.00", "4001", "19931", "0.60", "-0.66", "[8745.29 ; 9182.48]"],
        ["Forks", "1318.58", "1211.51", "998.00", "8", "19556", "4.74", "48.69", "[1258.55 ; 1378.61]"],
        ["Open Issues", "179.91", "337.52", "78.00", "0", "5959", "6.71", "76.95", "[163.18 ; 196.63]"],
        ["Size (KB)", "136916.26", "769547.30", "16923.00", "3", "20737671", "16.94", "373.41", "[98784.76 ; 175047.76]"]
    ]
    create_custom_table(t1_data)

    add_p(
        "Estimación puntual y proporciones: El promedio estimado de stars es de 8963.88. Con 95 % de confianza, la media poblacional se ubica "
        "entre 8745.29 y 9182.48. Además, 573 repositorios (36.57 %) superan las 10 000 estrellas (IC 95 %: [34.18 % ; 38.95 %]) y 130 repositorios "
        "(8.30 %) superan las 500 incidencias abiertas (IC 95 %: [6.93 % ; 9.67 %])."
    )

    # =========================================================================
    # 4. VERIFICACIÓN DE SUPUESTOS
    # =========================================================================
    add_h1("4. Verificación de supuestos: Normalidad y Homogeneidad de Varianzas")
    add_p(
        "Antes de aplicar las pruebas inferenciales se verificaron los supuestos de normalidad (Shapiro–Wilk) y homogeneidad de varianzas (Levene), "
        "conforme a las indicaciones del docente (α = 0.05). La Tabla 2 expone las pruebas de Shapiro–Wilk sobre la muestra completa y desglosadas por propietario."
    )

    add_p("Tabla 2. Pruebas de normalidad de Shapiro–Wilk para variables originales y transformadas.", bold=True, size=10, align=WD_ALIGN_PARAGRAPH.CENTER, color=NAVY)
    t2_data = [
        ["Variable", "Grupo / Muestra", "N", "Estadístico W", "p-valor", "¿Sigue Dist. Normal? (α = 0.05)"],
        ["Stars", "Muestra Global", "1567", "0.8989", "< 0.001", "No (p < 0.05, Rechaza H₀)"],
        ["Stars", "Organization", "878", "0.9076", "< 0.001", "No (p < 0.05, Rechaza H₀)"],
        ["Stars", "User", "689", "0.8861", "< 0.001", "No (p < 0.05, Rechaza H₀)"],
        ["Forks", "Muestra Global", "1567", "0.6878", "< 0.001", "No (p < 0.05, Rechaza H₀)"],
        ["Open Issues", "Muestra Global", "1567", "0.4757", "< 0.001", "No (p < 0.05, Rechaza H₀)"],
        ["ln_forks_1", "Organization", "878", "0.9774", "< 0.001", "No (TLC garantiza convergencia)"],
        ["ln_forks_1", "User", "689", "0.9972", "0.2974", "Sí (p ≥ 0.05, Acepta H₀)"],
        ["Diff (stars - forks)", "Diferencias Emparejadas", "1567", "0.9662", "< 0.001", "No (TLC garantiza convergencia)"]
    ]
    create_custom_table(t2_data)

    add_p(
        "Fundamentación metodológica: Aunque las variables brutas rechazan la normalidad estricta (p < 0.001), el Teorema del Límite Central (TLC) "
        "establece formalmente que para tamaños de muestra grandes (n ≥ 30, y en este estudio n = 1567 con subgrupos de n ≥ 405), la distribución muestral "
        "de la media converge a una normal. Por tanto, el uso de las pruebas paramétricas t y ANOVA queda matemáticamente justificado."
    )

    add_p("Tabla 3. Pruebas de homogeneidad de varianzas de Levene (Organization vs. User).", bold=True, size=10, align=WD_ALIGN_PARAGRAPH.CENTER, color=NAVY)
    t3_data = [
        ["Variable", "Estadístico Levene (F)", "gl1", "gl2", "p-valor", "¿Varianzas Homogéneas?", "Prueba Seleccionada"],
        ["Stars", "3.885", "1", "1565", "0.0489", "No (Heterocedasticidad)", "t de Welch"],
        ["Forks", "3.068", "1", "1565", "0.0801", "Sí (Homocedasticidad)", "t de Student"],
        ["ln_forks_1 (SPSS)", "7.352", "1", "1565", "0.0070", "No (Heterocedasticidad)", "t de Welch"],
        ["Open Issues", "40.191", "1", "1565", "< 0.001", "No (Heterocedasticidad)", "t de Welch"],
        ["Size (KB)", "0.066", "1", "1565", "0.7968", "Sí (Homocedasticidad)", "t de Student"]
    ]
    create_custom_table(t3_data)

    # =========================================================================
    # 5. PROCEDIMIENTO 1: PRUEBA DE UNA MUESTRA
    # =========================================================================
    add_h1("5. Procedimiento 1: Prueba de hipótesis para una muestra (stars)")
    add_p("Se desarrolla el procedimiento siguiendo rigurosamente la secuencia didáctica del ingeniero:")

    add_p("1. Formulación de hipótesis (Prueba bilateral, nivel de significancia α = 0.05):", bold=True)
    add_p("   • H₀: μ = 9000  (El promedio poblacional de estrellas es igual al valor de referencia de 9000).")
    add_p("   • H₁: μ ≠ 9000  (El promedio poblacional de estrellas es diferente de 9000).")

    add_p("2. Estadísticos descriptivos de la muestra (n = 1567):", bold=True)
    add_p("   • Media muestral: x̄ = 8963.88 estrellas.  • Desviación estándar: s = 4411.51 estrellas.")
    add_p("   • Error estándar de la media: SE = s / √n = 111.44 estrellas.  • Diferencia de medias: x̄ - μ₀ = -36.12.")

    add_p("3. Verificación del supuesto de normalidad:", bold=True)
    add_p("   • Shapiro–Wilk: W = 0.899 (gl = 1567, p < 0.001). Aunque p < 0.05, el Teorema del Límite Central (n = 1567 >> 30) asegura la distribución normal de la media muestral.")

    add_p("4. Aplicación de la prueba t para una muestra:", bold=True)
    add_p("   • Estadístico de prueba: t = (x̄ - μ₀) / SE = (8963.88 - 9000) / 111.44 = -0.324.")
    add_p("   • Grados de libertad: gl = n - 1 = 1566.")
    add_p("   • p-valor obtenido en SPSS: p = 0.746 (bilateral) / p = 0.373 (un factor).")
    add_p("   • Intervalo de confianza al 95 % de la diferencia: [-254.71 ; 182.48].")

    add_p("5. Toma de decisión estadística según la regla:", bold=True)
    add_p("   • Regla de decisión: Si p-valor ≥ α (0.05) → No se rechaza H₀. Si p-valor < α (0.05) → Se rechaza H₀.")
    add_p("   • Decisión: Como p-valor = 0.746 ≥ 0.05, NO SE RECHAZA LA HIPÓTESIS NULA (H₀).")

    add_p("6. Tamaño del efecto (d de Cohen y corrección de Hedges):", bold=True)
    add_p("   • d de Cohen = (x̄ - μ₀) / s = -0.008 (IC 95 %: [-0.058 ; 0.041]).  • Hedges g = -0.008. Magnitud: Efecto nulo / trivial.")

    add_p("7. Conclusión e interpretación en el contexto de Ingeniería de Sistemas:", bold=True)
    add_p(
        "   No existe evidencia estadística de que el promedio de estrellas de los repositorios de Python difiera del estándar de 9000. "
        "   Los proyectos analizados pertenecen a un estrato maduro cuya popularidad se alinea con precisión alrededor de 9000 estrellas. "
        "   (Complementariamente, frente al umbral histórico mínimo de 5000 estrellas, la prueba arroja t = 35.57, p < 0.001, d = 0.90, rechazándose H₀ con efecto grande)."
    )

    # Images for Section 5 (Compact 4.0 - 4.2 in)
    add_img("extracted_trabajo_media/image15.png", "Figura 1. Estadísticos descriptivos y pruebas de normalidad de stars en SPSS.")
    add_img("extracted_trabajo_media/image9.png", "Figura 2. Histograma de frecuencias de stars con parámetros muestrales (x̄ = 8963.88, s = 4411.51).")
    add_img("extracted_trabajo_media/image7.png", "Figura 3. Gráfico Q-Q normal de la variable stars.")
    add_img("extracted_trabajo_media/image11.png", "Figura 4. Diagrama de caja (Boxplot) de stars mostrando cuartiles y mediana.")
    add_img("extracted_trabajo_media/image5.png", "Figura 5. Resultados de la prueba t para una muestra (μ₀ = 9000) y tamaños de efecto en SPSS.")

    # =========================================================================
    # 6. PROCEDIMIENTO 2: DOS GRUPOS INDEPENDIENTES
    # =========================================================================
    add_h1("6. Procedimiento 2: Comparación de promedios entre dos grupos independientes (ln_forks_1)")
    add_p("Se contrasta el nivel de bifurcaciones transformadas (ln_forks_1) según el tipo de propietario (Organization vs. User):")

    add_p("1. Formulación de hipótesis (Prueba bilateral, nivel de significancia α = 0.05):", bold=True)
    add_p("   • H₀: μ_Org = μ_User  (El promedio de ln_forks_1 es igual entre organizaciones y usuarios individuales).")
    add_p("   • H₁: μ_Org ≠ μ_User  (El promedio de ln_forks_1 es diferente entre organizaciones y usuarios individuales).")

    add_p("2. Estadísticos descriptivos por grupo:", bold=True)
    add_p("   • Grupo Organization (n₁ = 878): Media x̄₁ = 6.9141, Desv. Est. s₁ = 0.7477, Error Est. SE₁ = 0.0252.")
    add_p("   • Grupo User (n₂ = 689): Media x̄₂ = 6.8701, Desv. Est. s₂ = 0.8030, Error Est. SE₂ = 0.0306.")
    add_p("   • Diferencia de medias: x̄₁ - x̄₂ = 0.04391 unidades logarítmicas.")

    add_p("3. Verificación de supuestos:", bold=True)
    add_p("   • Normalidad (Shapiro–Wilk): Organization W = 0.9774 (p < 0.001); User W = 0.9972 (p = 0.2974, cumple normalidad al 5 %). El grupo corporativo se ampara en el TLC (n₁ = 878 >> 30).")
    add_p("   • Homogeneidad de varianzas (Prueba de Levene): Estadístico F = 7.352, p = 0.007. Dado que p < 0.05, se RECHAZA la homocedasticidad (varianzas heterogéneas).")

    add_p("4. Selección y aplicación de la prueba inferencial: t de Welch:", bold=True)
    add_p("   • Ante varianzas desiguales, se aplica la prueba t de Welch con grados de libertad ajustados de Satterthwaite.")
    add_p("   • Estadístico t = 1.107.  • Grados de libertad ajustados: gl = 1425.171.")
    add_p("   • p-valor obtenido en SPSS: p = 0.268 (bilateral) / p = 0.134 (un factor).")
    add_p("   • Error estándar de la diferencia: SE_dif = 0.03966.  • IC 95 % de la diferencia: [-0.03388 ; 0.12171].")

    add_p("5. Toma de decisión estadística según la regla:", bold=True)
    add_p("   • Regla: p-valor ≥ α (0.05) → No se rechaza H₀. p-valor < α (0.05) → Se rechaza H₀.")
    add_p("   • Decisión: Como p-valor = 0.268 ≥ 0.05, NO SE RECHAZA LA HIPÓTESIS NULA (H₀).")

    add_p("6. Tamaño del efecto para dos grupos independientes:", bold=True)
    add_p("   • d de Cohen = 0.057 (IC 95 %: [-0.043 ; 0.157]).  • g de Hedges = 0.057.  • Delta de Glass = 0.055. Magnitud: Efecto trivial.")

    add_p("7. Conclusión e interpretación en el contexto del software:", bold=True)
    add_p(
        "   No existe diferencia estadísticamente significativa en el nivel de bifurcaciones transformadas entre proyectos de organizaciones "
        "   y usuarios independientes. La capacidad técnica de un código para ser derivado y reutilizado es independiente del respaldo corporativo."
    )

    # Images for Section 6
    add_img("extracted_trabajo_media/image2.png", "Figura 6. Prueba de Levene para igualdad de varianzas de ln_forks_1 (F = 7.352, p = 0.007).")
    add_img("extracted_trabajo_media/image12.png", "Figura 7. Estadísticas de grupo y Prueba t para muestras independientes (t de Welch = 1.107, p = 0.268) en SPSS.")
    add_img("extracted_trabajo_media/image8.png", "Figura 8. Tamaños de efecto para muestras independientes (d = 0.057, g = 0.057, Delta = 0.055).")
    add_img("extracted_reporte_media/image1.png", "Figura 9. Histograma de distribución de stars con curva normal para repositorios de Organizaciones (Reporte Estadístico).")

    # =========================================================================
    # 7. PROCEDIMIENTO 3: MUESTRAS RELACIONADAS
    # =========================================================================
    add_h1("7. Procedimiento 3: Comparación de promedios en muestras relacionadas (ln_stars_1 vs. ln_forks_1)")
    add_p("Se contrasta la popularidad pasiva frente a la adopción técnica sobre las MISMAS unidades de análisis (los 1567 repositorios):")

    add_p("1. Justificación metodológica del emparejamiento:", bold=True)
    add_p(
        "   Cada repositorio en GitHub registra simultáneamente dos métricas: ln_stars_1 y ln_forks_1. Al provenir del mismo artefacto de código, "
        "   comparten condiciones de calidad, dominio y autoría. El diseño emparejado elimina la variabilidad inter-repositorio y evalúa la discrepancia intra-proyecto."
    )

    add_p("2. Formulación de hipótesis (Prueba bilateral sobre la diferencia media μ_D, α = 0.05):", bold=True)
    add_p("   • H₀: μ_D = 0  (El promedio de las diferencias entre estrellas y forks es igual a cero; μ_stars = μ_forks).")
    add_p("   • H₁: μ_D ≠ 0  (El promedio de las diferencias difiere de cero; existe discrepancia significativa entre estrellas y forks).")

    add_p("3. Estadísticos descriptivos de las muestras emparejadas (n = 1567):", bold=True)
    add_p("   • ln_stars_1: Media = 8.977201, Desv. Est. = 0.503141, Error Est. = 0.012710.")
    add_p("   • ln_forks_1: Media = 6.894743, Desv. Est. = 0.772577, Error Est. = 0.019517.")
    add_p("   • Correlación de Pearson entre pares: r = 0.623 (p < 0.001). Correlación positiva moderada-fuerte que respalda el emparejamiento.")
    add_p("   • Diferencia emparejada (D = ln_stars_1 - ln_forks_1): Media D̄ = 2.082458, Desv. Est. s_D = 0.604453, Error Est. SE_D = 0.015270.")
    add_p("   • Intervalo de confianza al 95 % de la diferencia: [2.052507 ; 2.112409].")

    add_p("4. Verificación de normalidad de las diferencias:", bold=True)
    add_p("   • Shapiro–Wilk sobre D: W = 0.9662 (gl = 1567, p < 0.001). El TLC garantiza la normalidad asintótica de la media de diferencias D̄.")

    add_p("5. Aplicación de la prueba t para muestras relacionadas:", bold=True)
    add_p("   • Estadístico t = D̄ / SE_D = 2.082458 / 0.015270 = 136.379.")
    add_p("   • Grados de libertad: gl = n - 1 = 1566.")
    add_p("   • p-valor obtenido en SPSS: p < 0.001 (bilateral y un factor).")

    add_p("6. Toma de decisión estadística según la regla:", bold=True)
    add_p("   • Regla: Si p-valor < α (0.05) → Se rechaza H₀. Decisión: Como p-valor < 0.001 < 0.05, SE RECHAZA ROTUNDAMENTE H₀ a favor de H₁.")

    add_p("7. Tamaño del efecto para muestras emparejadas:", bold=True)
    add_p("   • d de Cohen = D̄ / s_D = 3.445 (IC 95 %: [3.315 ; 3.575]).  • g de Hedges = 3.444. Magnitud: Efecto extremadamente masivo (d >> 0.80).")

    add_p("8. Conclusión e interpretación en el contexto de Ingeniería de Sistemas:", bold=True)
    add_p(
        "   Existe una disparidad masiva y concluyente a favor de las estrellas sobre los forks. Una diferencia logarítmica de 2.0825 equivale "
        "   a e^(2.0825) ≈ 8.02 veces más estrellas que bifurcaciones por repositorio. Esto formaliza cuantitativamente que el respaldo comunitario pasivo "
        "   (un clic) es 8 veces más frecuente que la clonación activa para desarrollo derivado."
    )

    # Images for Section 7
    add_img("extracted_trabajo_media/image16.png", "Figura 10. Estadísticas de muestras emparejadas, correlación (r = 0.623) y Prueba t de muestras relacionadas (t = 136.379, p < 0.001) en SPSS.")
    add_img("extracted_trabajo_media/image1.png", "Figura 11. Ayuda curada y tamaños de efecto de muestras emparejadas (d de Cohen = 3.445, g = 3.444).")

    # =========================================================================
    # 8. PROCEDIMIENTO 4: ANOVA PARA TRES GRUPOS
    # =========================================================================
    add_h1("8. Procedimiento 4: Análisis de Varianza (ANOVA) para tres grupos independientes (ln_open_issues_1)")
    add_p("Se contrasta la carga de incidencias abiertas transformadas según el marco de licenciamiento (MIT, Apache-2.0 y GPL-3.0):")

    add_p("1. Formulación de hipótesis (Nivel de significancia α = 0.05):", bold=True)
    add_p("   • H₀: μ_MIT = μ_Apache = μ_GPL  (Los promedios de ln_open_issues_1 son iguales entre las tres licencias).")
    add_p("   • H₁: Al menos un promedio poblacional de ln_open_issues_1 es diferente.")

    add_p("2. Estadísticos descriptivos por grupo de licencia (N = 1063):", bold=True)
    add_p("   • Licencia MIT (lic3 = 1.00; n₁ = 546): Media x̄₁ = 4.0227, Desv. Est. s₁ = 1.4769, Error Est. = 0.0632, IC 95%: [3.8985 ; 4.1469].")
    add_p("   • Licencia Apache-2.0 (lic3 = 2.00; n₂ = 405): Media x̄₂ = 4.5735, Desv. Est. s₂ = 1.3876, Error Est. = 0.0689, IC 95%: [4.4379 ; 4.7090].")
    add_p("   • Licencia GPL-3.0 (lic3 = 3.00; n₃ = 112): Media x̄₃ = 4.3647, Desv. Est. s₃ = 1.3920, Error Est. = 0.1315, IC 95%: [4.1041 ; 4.6254].")
    add_p("   • Total analizado (N = 1063): Media = 4.2686, Desv. Est. = 1.4567, Error Est. = 0.0447, IC 95%: [4.1809 ; 4.3563].")

    add_p("3. Verificación de supuestos:", bold=True)
    add_p("   • Normalidad (Shapiro–Wilk): MIT W = 0.9817 (p < 0.001); Apache-2.0 W = 0.9760 (p < 0.001); GPL-3.0 W = 0.9801 (p = 0.0930, cumple normalidad al 5 %). Celdas grandes (n ≥ 112) garantizan robustez del ANOVA.")
    add_p("   • Homogeneidad de varianzas (Prueba de Levene): Estadístico F = 1.109, gl1 = 2, gl2 = 1060, p = 0.330. Como p = 0.330 ≥ 0.05, NO SE RECHAZA la homocedasticidad. Las varianzas poblacionales son homogéneas.")

    add_p("4. Aplicación del ANOVA de un factor y prueba de Welch:", bold=True)
    add_p("   • Suma de cuadrados Entre Grupos: SS_entre = 71.698, gl = 2, Media cuadrática MS_entre = 35.849.")
    add_p("   • Suma de cuadrados Dentro de Grupos: SS_dentro = 2181.693, gl = 1060, Media cuadrática MS_dentro = 2.058.")
    add_p("   • Suma de cuadrados Total: SS_total = 2253.391, gl = 1062.")
    add_p("   • Estadístico F calculado: F = MS_entre / MS_dentro = 35.849 / 2.058 = 17.418.")
    add_p("   • p-valor obtenido en SPSS: p < 0.001.")
    add_p("   • Prueba robusta de Welch: F_Welch = 17.511 (gl1 = 2, gl2 = 314.565, p < 0.001), corroborando idéntico resultado.")

    add_p("5. Toma de decisión estadística según la regla:", bold=True)
    add_p("   • Regla: p-valor < α (0.05) → Se rechaza H₀. Decisión: Como p-valor < 0.001 < 0.05, SE RECHAZA H₀ y se acepta H₁.")

    add_p("6. Tamaño del efecto para el ANOVA:", bold=True)
    add_p("   • Eta cuadrado (η²) = SS_entre / SS_total = 71.698 / 2253.391 = 0.032 (IC 95 %: [0.014 ; 0.054]).  • Omega cuadrado (ω²) = 0.030. Magnitud: Pequeña a moderada (la licencia explica el 3.2 % de la varianza en issues).")

    add_p("7. Conclusión e interpretación en el contexto de software:", bold=True)
    add_p(
        "   Existen diferencias estadísticamente significativas en el promedio de incidencias abiertas transformadas entre los tipos de licencia. "
        "   Sin embargo, el ANOVA global no identifica entre qué grupos se presentan las diferencias; por ende, es obligatorio continuar con las pruebas Post Hoc."
    )

    # Images for Section 8
    add_img("extracted_trabajo_media/image10.png", "Figura 12. Descriptivos por tipo de licencia y Prueba de Homogeneidad de Varianzas de Levene (F = 1.109, p = 0.330) en SPSS.")
    add_img("extracted_trabajo_media/image3.png", "Figura 13. Tabla ANOVA de un factor (F = 17.418, p < 0.001), tamaños de efecto y prueba de Welch en SPSS.")

    # =========================================================================
    # 9. PROCEDIMIENTO 5: PRUEBAS POST HOC
    # =========================================================================
    add_h1("9. Procedimiento 5: Pruebas Post Hoc de comparaciones múltiples (HSD Tukey)")
    add_p("Siguiendo la directriz del docente: Al cumplirse la homogeneidad de varianzas en Levene (p = 0.330 ≥ 0.05), se utiliza la prueba HSD de Tukey:")

    add_p("1. Criterio de selección justificado:", bold=True)
    add_p("   • Varianzas homogéneas (Levene p ≥ 0.05) → HSD de Tukey. (Si hubiesen sido heterogéneas, se habría empleado Games–Howell).")

    add_p("2. Comparaciones múltiples por pares (HSD Tukey, α = 0.05):", bold=True)
    add_p("   • Par 1: MIT vs. Apache-2.0:")
    add_p("     Diferencia de medias (x̄_MIT - x̄_Apache) = -0.550792, Error Est. = 0.094083, p < 0.001, IC 95%: [-0.771605 ; -0.329979].")
    add_p("     Decisión: DIFERENCIA ALTAMENTE SIGNIFICATIVA (p < 0.05, marcado con * en SPSS). Los proyectos Apache-2.0 tienen significativamente más issues abiertos.")
    add_p("   • Par 2: MIT vs. GPL-3.0:")
    add_p("     Diferencia de medias (x̄_MIT - x̄_GPL) = -0.342007, Error Est. = 0.148817, p = 0.056, IC 95%: [-0.691280 ; 0.007266].")
    add_p("     Decisión: NO SIGNIFICATIVA al nivel 0.05 (p = 0.056 ≥ 0.05, tendencia marginal).")
    add_p("   • Par 3: Apache-2.0 vs. GPL-3.0:")
    add_p("     Diferencia de medias (x̄_Apache - x̄_GPL) = 0.208785, Error Est. = 0.153163, p = 0.361, IC 95%: [-0.150688 ; 0.568258].")
    add_p("     Decisión: NO SIGNIFICATIVA (p = 0.361 ≥ 0.05). No hay diferencia entre Apache-2.0 y GPL-3.0.")

    add_p("3. Análisis de subconjuntos homogéneos (HSD Tukey, media armónica n_h = 226.771):", bold=True)
    add_p("   • Subconjunto 1 (Baja carga de issues): Licencia MIT (media = 4.022699; Sig. = 1.000).")
    add_p("   • Subconjunto 2 (Alta carga de issues): Licencia GPL-3.0 (media = 4.364706) y Licencia Apache-2.0 (media = 4.573491; Sig. = 0.268).")

    add_p("4. Conclusión e interpretación contextual en Ingeniería de Software:", bold=True)
    add_p(
        "   La disparidad detectada por el ANOVA se concentra de manera exclusiva entre Apache-2.0 y MIT. "
        "   Los proyectos Apache-2.0 corresponden comúnmente a infraestructuras de gobernanza corporativa compleja y auditorías de código empresarial "
        "   que generan prolongados ciclos de tickets abiertos. En contraste, MIT fomenta soluciones minimalistas y ligeras con rápida resolución de incidencias. "
        "   GPL-3.0 se agrupa con Apache-2.0 debido a los debates comunitarios extensos sobre compatibilidad de licencias."
    )

    # Images for Section 9
    add_img("extracted_trabajo_media/image13.png", "Figura 14. Tabla de Comparaciones Múltiples Post Hoc HSD de Tukey para ln_open_issues_1 en SPSS.")
    add_img("extracted_trabajo_media/image4.png", "Figura 15. Subconjuntos homogéneos de Tukey agrupando a MIT en el Subconjunto 1 y a GPL-3.0 y Apache-2.0 en el Subconjunto 2.")

    # =========================================================================
    # 10. CONCLUSIONES GENERALES
    # =========================================================================
    add_h1("10. Conclusiones generales")
    add_p("A partir de la evidencia estadística inferencial obtenida sobre los 1567 repositorios de software libre en Python:")
    add_p("1. Popularidad de una muestra (Q1): El promedio de estrellas de los repositorios de Python no difiere estadísticamente de 9000 (t = -0.324, p = 0.746, d = -0.008), respaldando que la muestra analizada se consolida alrededor de este umbral de impacto. Frente a 5000 estrellas, los proyectos superan holgadamente dicha cota (p < 0.001, d = 0.90).")
    add_p("2. Comparación entre grupos independientes (Q2): No existen diferencias estadísticamente significativas en el nivel de bifurcaciones transformadas entre Organizaciones y Usuarios individuales (t = 1.107, p = 0.268, d = 0.057). La reutilización de código se fundamenta en el valor intrínseco de la solución y no en la naturaleza organizativa del autor.")
    add_p("3. Muestras relacionadas intra-repositorio (Q3): Se constató una diferencia masiva a favor de las estrellas sobre las bifurcaciones dentro de los mismos proyectos (t = 136.379, p < 0.001, d = 3.445). En la dinámica de código abierto, existe una relación promedio de 8 estrellas por cada bifurcación, confirmando la asimetría entre apoyo pasivo y contribución técnica activa.")
    add_p("4. ANOVA de un factor (Q4): El modelo de licenciamiento influye significativamente en la acumulación de incidencias abiertas transformadas (F = 17.418, p < 0.001, η² = 0.032). Se rechaza formalmente la igualdad de medias poblacionales entre licencias.")
    add_p("5. Pruebas Post Hoc de Tukey (Q5): Habiéndose verificado varianzas homogéneas en Levene (p = 0.330), la prueba HSD de Tukey demostró que la diferencia significativa radica exclusivamente entre Apache-2.0 y MIT (p < 0.001). Apache-2.0 y GPL-3.0 conforman un clúster de mayor complejidad de mantenimiento frente a la agilidad de MIT.")
    add_p("6. Robustez metodológica: A pesar del rechazo de la normalidad estricta en Shapiro–Wilk (p < 0.001), los contrastes paramétricos mantuvieron plena validez sustentados en el Teorema del Límite Central (n = 1567 >> 30), transformaciones logarítmicas estabilizadoras y correcciones de varianzas heterogéneas.")

    # =========================================================================
    # 11. REFERENCIAS Y FUENTES DE DATOS
    # =========================================================================
    add_h1("11. Referencias y fuentes de datos")
    refs = [
        "Cohen, J. (1988). Statistical Power Analysis for the Behavioral Sciences (2nd ed.). Lawrence Erlbaum Associates.",
        "Hedges, L. V. (1981). Distribution theory for Glass's estimator of effect size and related estimators. Journal of Educational Statistics, 6(2), 107–128.",
        "IBM Corp. (2023). IBM SPSS Statistics for Windows, Version 29.0. Armonk, NY: IBM Corp.",
        "Kaggle Datasets. (2024). GitHub Repositories – Heterogeneous Data. Extraído de: https://www.kaggle.com/datasets",
        "Montgomery, D. C., & Runger, G. C. (2018). Probabilidad y estadística aplicadas a la ingeniería (5.ª ed.). Limusa Wiley.",
        "SciPy Community. (2024). Statistical functions (scipy.stats). https://docs.scipy.org/doc/scipy/reference/stats.html",
        "Triola, M. F. (2018). Estadística (12.ª ed.). Pearson Educación."
    ]

    for ref in refs:
        p_r = doc.add_paragraph()
        p_r.paragraph_format.space_before = Pt(1)
        p_r.paragraph_format.space_after = Pt(3)
        p_r.paragraph_format.left_indent = Inches(0.4)
        p_r.paragraph_format.first_line_indent = Inches(-0.4)
        p_r.paragraph_format.line_spacing = 1.15
        r_r = p_r.add_run(ref)
        r_r.font.name = 'Times New Roman'
        r_r.font.size = Pt(10)

    # Save document
    output_filename = "PEREZ_AROCUTIPA_LLICA_APLICACION_ESTADISTICA_INFERENCIAL.docx"
    doc.save(output_filename)
    print(f"Streamlined document successfully created and saved as: {output_filename}")
    try:
        doc.save("INFORME_FINAL_APLICACION_DE_ESTADISTICA_INFERENCIAL.docx")
        print("Also updated INFORME_FINAL_APLICACION_DE_ESTADISTICA_INFERENCIAL.docx")
    except Exception as e:
        print("Note: Previous file was locked by viewer; primary file saved as PEREZ_AROCUTIPA_LLICA_APLICACION_ESTADISTICA_INFERENCIAL.docx")

if __name__ == '__main__':
    generate_streamlined_doc()
