# -*- coding: utf-8 -*-
import pyreadstat
import pandas as pd
import os
import docx

print("=== 1. INSPECTING datos.sav ===")
df_sav, meta = pyreadstat.read_sav('datos.sav')
print(f"datos.sav shape: {df_sav.shape[0]} rows x {df_sav.shape[1]} columns")
print(f"Number of variables in datos.sav: {len(meta.column_names)}")
print("Variable names in datos.sav:")
for idx, col in enumerate(meta.column_names):
    lbl = meta.column_labels[idx] if idx < len(meta.column_labels) else ""
    print(f"  [{idx+1}] {col}: label='{lbl}', type={df_sav[col].dtype}")

print("\n=== 2. INSPECTING Heterogeneous_Data .csv ===")
csv_path = r'c:\Users\PC\Desktop\Estats\archive\Heterogeneous_Data .csv'
if os.path.exists(csv_path):
    df_sample = pd.read_csv(csv_path, nrows=5)
    print(f"Number of columns in Heterogeneous_Data .csv: {len(df_sample.columns)}")
    print("CSV column names:")
    for idx, col in enumerate(df_sample.columns):
        print(f"  [{idx+1}] {col}")

print("\n=== 3. SEARCHING FOR '[COMPLETAR' IN DOCX FILES ===")
for fname in os.listdir('.'):
    if fname.endswith('.docx') and not fname.startswith('~$'):
        try:
            doc = docx.Document(fname)
            found = []
            for p_idx, p in enumerate(doc.paragraphs):
                if 'COMPLETAR' in p.text or 'TOTAL DE VARIABLES' in p.text or 'marcador' in p.text:
                    found.append((p_idx, p.text))
            for t_idx, t in enumerate(doc.tables):
                for r_idx, r in enumerate(t.rows):
                    for c_idx, c in enumerate(r.cells):
                        if 'COMPLETAR' in c.text or 'TOTAL DE VARIABLES' in c.text:
                            found.append((f"table {t_idx} r{r_idx} c{c_idx}", c.text))
            if found:
                print(f"Found in {fname}:")
                for item in found:
                    print(f"   {item[0]}: {item[1]}")
        except Exception as e:
            print(f"Error reading {fname}: {e}")
