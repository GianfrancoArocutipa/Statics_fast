# -*- coding: utf-8 -*-
"""
Script para generar el documento resuelto de la Segunda Práctica Calificada
Preserva el 100% del formato original (márgenes, encabezado, pie, banner negro, tipografías y estilos)
e inserta la resolución completa de los 5 problemas con imágenes SPSS oficiales.
"""

import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.parts.document import DocumentPart

# Parchear next_id para buscar IDs en todas las partes del paquete y evitar colisiones con formas del encabezado/pie
def _custom_next_id(self):
    used_ids = set()
    for part in self.package.parts:
        if hasattr(part, '_element') and part._element is not None:
            for s in part._element.xpath('//@id'):
                if s.isdigit():
                    val = int(s)
                    if val < 1000000:
                        used_ids.add(val)
    return (max(used_ids) + 1) if used_ids else 2000

DocumentPart.next_id = property(_custom_next_id)

base_doc_path = os.path.join("prueba", "Enunciado Segunda practica calificada inferencia estadistica 2026-II.docx")
output_doc_path = os.path.join("prueba", "Segunda_Practica_Calificada_Inferencia_Estadistica_2026-II_Resuelto.docx")
images_dir = os.path.join("prueba", "spss_images")

doc = Document(base_doc_path)

# Estilos auxiliares
def add_styled_paragraph(doc, text="", style='List Paragraph', align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=2, space_before=2):
    p = doc.add_paragraph(style=style)
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    if text:
        r = p.add_run(text)
        r.font.name = 'Calibri'
        r.font.size = Pt(11)
        r.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
    return p

def add_heading_paragraph(doc, text, style='List Paragraph', bold=True, space_before=4, space_after=2):
    p = doc.add_paragraph(style=style)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(text)
    r.font.name = 'Calibri'
    r.font.size = Pt(11)
    r.font.bold = bold
    r.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
    return p

def add_centered_image(doc, img_name, width_inches=5.2, space_before=4, space_after=4):
    img_path = os.path.join(images_dir, img_name)
    if not os.path.exists(img_path):
        print(f"Alerta: Imagen {img_path} no encontrada.")
        return None
    p = doc.add_paragraph(style='Normal')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    r = p.add_run()
    r.add_picture(img_path, width=Inches(width_inches))
    return p

# Limpiar párrafos vacíos o enunciados posteriores al banner negro para reconstruir con las soluciones
# En el doc original:
# P0-P7 son el encabezado y banner negro.
# P9: Enunciado 1
# P11: Enunciado 2
# P13-17: Enunciado 3
# P20: Enunciado 4
# P22-28: Enunciado 5
# P34: Tacna, octubre de 2026

# Guardamos los textos exactos de los enunciados del doc original para preservar el texto idéntico
p_texts = [p.text for p in doc.paragraphs]

# Reutilizaremos el documento base, conservando encabezado, banner negro y secciones
# Vaciaremos los párrafos a partir de P8 para volver a ensamblar con sus resoluciones
paragraphs_to_keep = 8  # P0 a P7
for i in range(len(doc.paragraphs) - 1, paragraphs_to_keep - 1, -1):
    p = doc.paragraphs[i]._p
    p.getparent().remove(p)

print(f"Documento podado a {len(doc.paragraphs)} párrafos base.")

# ========================================================
# PROBLEMA 1
# ========================================================
p_enunc1 = add_styled_paragraph(
    doc,
    "Una empresa de desarrollo de software desea determinar si los incidentes reportados por los usuarios de una aplicación web se distribuyen uniformemente entre cuatro categorías: errores de autenticación, fallas de conexión, problemas de rendimiento y errores de interfaz. Según el criterio técnico del equipo de desarrollo, se espera que las cuatro categorías presenten la misma proporción de incidentes. Para verificar este supuesto, se seleccionó una muestra aleatoria de 120 incidentes registrados durante el último mes, de los cuales 42 correspondieron a errores de autenticación, 25 a fallas de conexión, 33 a problemas de rendimiento y 20 a errores de interfaz. ¿Existe evidencia estadística suficiente para concluir que los incidentes reportados no se distribuyen uniformemente entre las cuatro categorías? Utilice un nivel de significancia del 3%. (4 ptos.)",
    style='List Paragraph',
    space_before=6,
    space_after=4
)

# Resolución Problema 1
add_heading_paragraph(doc, "Planteamiento de Hipótesis")
add_styled_paragraph(doc, "H0: Las proporciones observadas de incidentes reportados en la muestra son iguales a las proporciones esperadas de la población (se distribuyen uniformemente entre las cuatro categorías).")
add_styled_paragraph(doc, "p(Autenticación) = 25%, p(Conexión) = 25%, p(Rendimiento) = 25%, p(Interfaz) = 25%")
add_styled_paragraph(doc, "H1: Las proporciones observadas de incidentes reportados en la muestra son diferentes de las proporciones esperadas de la población (no se distribuyen uniformemente entre las cuatro categorías, al menos una de las proporciones difiere).")

add_heading_paragraph(doc, "Determinar el nivel de Significancia: 3% (0.03)")

add_heading_paragraph(doc, "Tabla de distribución de frecuencias")
add_centered_image(doc, "p1_frecuencias.png", width_inches=4.8)

add_heading_paragraph(doc, "Seleccionar el estadístico de prueba: Prueba no paramétrica Chi-cuadrado de bondad de ajuste")
add_centered_image(doc, "p1_estadisticos_prueba.png", width_inches=4.2)

add_heading_paragraph(doc, "Contrastación de Hipótesis")
add_styled_paragraph(doc, "Dado que p-valor (0.026) es menor que el nivel de significancia (0.03), es posible que se rechace la H0 (se rechaza la hipótesis nula).")

add_heading_paragraph(doc, "Conclusión")
add_styled_paragraph(doc, "A un margen del 3% se concluye que, los incidentes reportados no se distribuyen uniformemente entre las cuatro categorías.")

# ========================================================
# PROBLEMA 2
# ========================================================
add_styled_paragraph(
    doc,
    "Una empresa de desarrollo de software desea determinar si existen diferencias significativas en los tiempos de ejecución de una aplicación web al utilizar dos algoritmos de búsqueda de información. Para ello, se realizaron mediciones independientes del tiempo de respuesta, expresado en milisegundos, utilizando el algoritmo A en 10 ejecuciones y el algoritmo B en otras 11 ejecuciones. Los tiempos registrados para el algoritmo A fueron: 120, 135, 128, 142, 130, 125, 138, 132, 145 y 127 milisegundos; mientras que para el algoritmo B fueron: 110, 118, 115, 122, 108, 120, 112, 116, 125, 119 y 114 milisegundos. ¿Existe evidencia estadística suficiente para concluir que los tiempos de ejecución difieren significativamente entre ambos algoritmos? Considere un nivel de significancia del 1%. (4 ptos.)",
    style='List Paragraph',
    space_before=10,
    space_after=4
)

# Resolución Problema 2
add_heading_paragraph(doc, "Planteamiento de Hipótesis")
add_styled_paragraph(doc, "H0: No existe diferencia significativa en los tiempos de ejecución de la aplicación web entre el Algoritmo A y el Algoritmo B (los promedios son iguales: μA = μB).")
add_styled_paragraph(doc, "H1: Existe diferencia significativa en los tiempos de ejecución de la aplicación web entre el Algoritmo A y el Algoritmo B (los promedios difieren: μA ≠ μB).")

add_heading_paragraph(doc, "Determinar el nivel de Significancia: 1% (0.01)")

add_heading_paragraph(doc, "Tabla de estadísticos descriptivos (Informe)")
add_centered_image(doc, "p2_informe.png", width_inches=4.6)

add_heading_paragraph(doc, "Prueba de Distribución normal")
add_styled_paragraph(doc, "H0: Los tiempos de ejecución del Algoritmo A y del Algoritmo B presentan distribución normal")
add_styled_paragraph(doc, "H1: Los tiempos de ejecución de al menos un algoritmo no presentan distribución normal")
add_centered_image(doc, "p2_normalidad.png", width_inches=5.4)
add_styled_paragraph(doc, "Dado que los p-valores de Shapiro-Wilk para el Algoritmo A (0.957) y para el Algoritmo B (0.997) son mayores que el nivel de significancia (0.01), no se rechaza la hipótesis nula.")
add_styled_paragraph(doc, "Conclusión: A un margen de error del 1% se concluye que, los datos de ambos algoritmos presentan Distribución normal.")

add_heading_paragraph(doc, "Seleccionar el estadístico de prueba: Prueba paramétrica T de Student para muestras independientes")
add_styled_paragraph(doc, "Evaluación de homogeneidad de varianzas: La prueba de Levene reporta F = 1.855 con p-valor = 0.189 > 0.01, por lo que se asumen varianzas iguales.")
add_centered_image(doc, "p2_t_test.png", width_inches=5.8)
add_centered_image(doc, "p2_mann_whitney.png", width_inches=4.0)

add_heading_paragraph(doc, "Contrastación de Hipótesis")
add_styled_paragraph(doc, "Dado que p-valor (0.000) es menor que el nivel de significancia (0.01), es posible que se rechace la H0 (se rechaza la hipótesis nula).")

add_heading_paragraph(doc, "Conclusión")
add_styled_paragraph(doc, "A un margen del 1% se concluye que, los tiempos de ejecución difieren significativamente entre ambos algoritmos.")

# ========================================================
# PROBLEMA 3
# ========================================================
add_styled_paragraph(
    doc,
    "Un especialista en ciberseguridad desea determinar si los intentos de acceso a un sistema de información presentan un comportamiento aleatorio durante un periodo de monitoreo. Para ello, se registraron 30 intentos consecutivos de acceso, clasificándolos como acceso exitoso (E) o acceso fallido (F). La secuencia obtenida fue la siguiente:",
    style='List Paragraph',
    space_before=10,
    space_after=2
)
p_seq = add_styled_paragraph(
    doc,
    "E, E, F, E, F, F, E, E, F, E, F, F, E, F, E, E, F, F, E, F, E, E, F, E, F, F, E, E, F, E",
    style='List Paragraph',
    align=WD_ALIGN_PARAGRAPH.CENTER,
    space_before=2,
    space_after=2
)
p_seq.runs[0].font.bold = True

add_styled_paragraph(
    doc,
    "¿Existe evidencia estadística suficiente para concluir que la secuencia de intentos de acceso al sistema no presenta un comportamiento aleatorio? Considere un nivel de significancia del 5%. (4 ptos.)",
    style='List Paragraph',
    space_before=2,
    space_after=4
)

# Resolución Problema 3
add_heading_paragraph(doc, "Planteamiento de Hipótesis")
add_styled_paragraph(doc, "H0: No Existe un patrón o tendencia en la secuencia de intentos de acceso al sistema a lo largo del periodo de monitoreo (comportamiento aleatorio).")
add_styled_paragraph(doc, "H1: Existe un patrón o tendencia en la secuencia de intentos de acceso al sistema a lo largo del periodo de monitoreo (no presenta comportamiento aleatorio).")

add_heading_paragraph(doc, "Determinar el nivel de Significancia: 5% (0.05)")

add_heading_paragraph(doc, "Seleccionar el estadístico de prueba: Prueba no paramétrica Rachas")
add_centered_image(doc, "p3_rachas.png", width_inches=4.2)

add_heading_paragraph(doc, "Contrastación de Hipótesis")
add_styled_paragraph(doc, "Dado que p-valor (0.059) es mayor que el nivel de significancia (0.05), es probable que no se rechace la H0 (no se rechaza la hipótesis nula).")

add_heading_paragraph(doc, "Conclusión")
add_styled_paragraph(doc, "A un margen del 5% se concluye que, No Existe un patrón o tendencia en la secuencia de intentos de acceso al sistema a lo largo del periodo de monitoreo (la secuencia presenta un comportamiento aleatorio).")

# ========================================================
# PROBLEMA 4
# ========================================================
add_styled_paragraph(
    doc,
    "Un ingeniero de sistemas desea evaluar si la proporción de transacciones procesadas exitosamente por una nueva plataforma de pagos electrónicos cumple con el estándar de calidad establecido por la empresa. Según los registros históricos, se espera que el 90% de las transacciones se procesen correctamente. Para verificar este supuesto, se seleccionó una muestra aleatoria de 50 transacciones, de las cuales 40 fueron procesadas exitosamente y 10 presentaron errores durante su ejecución. ¿Existe evidencia estadística suficiente para concluir que la proporción de transacciones exitosas es inferior al estándar establecido del 90%? Considere un nivel de significancia del 5%. (4 ptos.)",
    style='List Paragraph',
    space_before=10,
    space_after=4
)

# Resolución Problema 4
add_heading_paragraph(doc, "Planteamiento de Hipótesis")
add_styled_paragraph(doc, "H0: La proporción de transacciones procesadas exitosamente no es inferior al estándar establecido del 90% (p ≥ 90%).")
add_styled_paragraph(doc, "H1: La proporción de transacciones procesadas exitosamente es inferior al estándar establecido del 90% (p < 90%).")

add_heading_paragraph(doc, "Determinar el nivel de Significancia: 5% (0.05)")

add_heading_paragraph(doc, "Tabla de distribución de frecuencias y Prueba binomial")
add_centered_image(doc, "p4_binomial.png", width_inches=5.2)

add_heading_paragraph(doc, "Seleccionar el estadístico de prueba: Prueba no paramétrica Binomial")

add_heading_paragraph(doc, "Contrastación de Hipótesis")
add_styled_paragraph(doc, "Dado que p-valor exacto unilateral (0.025) es menor que el nivel de significancia (0.05), es posible que se rechace la H0 (se rechaza la hipótesis nula).")

add_heading_paragraph(doc, "Conclusión")
add_styled_paragraph(doc, "A un margen del 5% se concluye que, la proporción de transacciones exitosas es inferior al estándar establecido del 90%.")

# ========================================================
# PROBLEMA 5
# ========================================================
add_styled_paragraph(
    doc,
    "Un ingeniero de sistemas desea determinar si existen diferencias significativas en los tiempos de respuesta de tres servicios de almacenamiento en la nube utilizados para recuperar archivos de una aplicación empresarial. Para ello, se realizaron mediciones independientes en cada servicio, registrando los tiempos de respuesta en milisegundos. Los tiempos registrados son:",
    style='List Paragraph',
    space_before=10,
    space_after=2
)
add_styled_paragraph(doc, "Servicio A: 95, 102, 98, 105, 100, 97, 103, 99, 106, 101 y 96 milisegundos;", style='List Paragraph', space_before=1, space_after=1)
add_styled_paragraph(doc, "Servicio B: 115, 122, 118, 125, 120, 117, 123, 119, 126, 121, 116, 124 y 128 milisegundos;", style='List Paragraph', space_before=1, space_after=1)
add_styled_paragraph(doc, "servicio C: 135, 142, 138, 145, 140, 137, 143, 139, 146, 141, 136, 144, 148, 134 y 147 milisegundos.", style='List Paragraph', space_before=1, space_after=2)
add_styled_paragraph(
    doc,
    "¿Existe evidencia estadística suficiente para concluir que los tiempos de respuesta difieren significativamente entre los tres servicios de almacenamiento en la nube? Considere un nivel de significancia del 3%. (4 ptos.)",
    style='List Paragraph',
    space_before=2,
    space_after=4
)

# Resolución Problema 5 (7 Pasos Oficiales UPT)
add_heading_paragraph(doc, "Paso 1: Formación de hipótesis")
add_styled_paragraph(doc, "H0: Las Promedios de los grupos son iguales (μA = μB = μC)")
add_styled_paragraph(doc, "H1: Por lo menos los Promedios de 2 grupos son diferentes (μi ≠ μj)")

add_heading_paragraph(doc, "Paso 2: Determinar el nivel de significancia: 3% (0.03)")

add_heading_paragraph(doc, "Paso 3: Hallar estadísticos descriptivos")
add_centered_image(doc, "p5_informe.png", width_inches=4.8)

add_heading_paragraph(doc, "Paso 4: Prueba de Distribución normal")
add_styled_paragraph(doc, "H0: Los datos de los tres servicios presentan distribución normal")
add_styled_paragraph(doc, "H1: Por lo menos los datos de un servicio no presentan distribución normal")
add_centered_image(doc, "p5_normalidad.png", width_inches=5.4)
add_styled_paragraph(doc, "Dado que los p-valores de los tres servicios (Servicio A: 0.885, Servicio B: 0.954, Servicio C: 0.755) son mayores que el nivel de significancia (0.03), no se rechaza la hipótesis nula.")
add_styled_paragraph(doc, "Conclusión: A un margen de error del 3% se concluye que, los datos de los grupos presentan Distribución normal.")

add_heading_paragraph(doc, "Paso 5: Elegir la prueba estadística")
add_styled_paragraph(doc, "Dado que los datos presentan distribución normal, se evalúa la homogeneidad de varianzas y se aplica el Análisis de Varianza (ANOVA de un factor, prueba paramétrica).")
add_styled_paragraph(doc, "Prueba de homogeneidad de varianzas (Levene): Dado que el p-valor (0.625) es mayor que el nivel de significancia (0.03), se asume homogeneidad de varianzas.")
add_centered_image(doc, "p5_levene.png", width_inches=4.8)
add_centered_image(doc, "p5_anova.png", width_inches=5.2)
add_styled_paragraph(doc, "Prueba Post Hoc de Tukey:")
add_centered_image(doc, "p5_posthoc.png", width_inches=5.2)
add_styled_paragraph(doc, "*(Alternativamente, si se aplicara la prueba no paramétrica Kruskal-Wallis: H = 33.600, p-valor = 0.000).*")
add_centered_image(doc, "p5_kruskal.png", width_inches=3.8)

add_heading_paragraph(doc, "Paso 6: Contrastación de Hipótesis")
add_styled_paragraph(doc, "Dado que el p-valor de ANOVA (0.000) es menor que el nivel de significancia (0.03), se rechaza la hipótesis nula.")

add_heading_paragraph(doc, "Paso 7: Conclusión")
add_styled_paragraph(doc, "A un margen de error del 3% se concluye que, por lo menos los tiempos de respuesta de 2 servicios de almacenamiento en la nube son diferentes.")

# Pie de documento
p_pie = doc.add_paragraph(style='Normal')
p_pie.alignment = WD_ALIGN_PARAGRAPH.RIGHT
p_pie.paragraph_format.space_before = Pt(16)
r_pie = p_pie.add_run("Tacna, octubre de 2026")
r_pie.font.name = 'Calibri'
r_pie.font.size = Pt(10)
r_pie.font.italic = True
r_pie.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

doc.save(output_doc_path)
print(f"Documento resuelto guardado exitosamente en: {output_doc_path}")
