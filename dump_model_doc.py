import docx

doc = docx.Document('PEREZ _AROCUTIPA_LLICA - APLICACIÓN DE ESTADÍSTICA INFERENCIAL CON DATOS REALES.docx')

with open('model_doc_full.txt', 'w', encoding='utf-8') as f:
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip():
            f.write(f"[{i}] {p.text}\n")
            
    f.write("\n\n=== TABLES ===\n")
    for t_idx, table in enumerate(doc.tables):
        f.write(f"\n--- TABLE {t_idx} ---\n")
        for r_idx, row in enumerate(table.rows):
            cells = [c.text.strip().replace('\n', ' ') for c in row.cells]
            f.write(f"Row {r_idx}: {' | '.join(cells)}\n")

print("Dumped model doc full text and tables.")
