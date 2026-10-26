import docx
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = docx.Document("Ejercicios Pruebas de chi cuadrado.docx")
target_p = None
for p in doc.paragraphs:
    if "En una empresa de desarrollo de software" in p.text:
        target_p = p
        break

if target_p:
    p_img = target_p.insert_paragraph_before("", style="List Paragraph")
    p_img.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run = p_img.add_run()
    run.add_picture("generated_spss_images/spss_chi_freq_so.png", width=Inches(3.5))
    doc.save("test_insert_img.docx")
    print("Image insertion before target succeeded!")
