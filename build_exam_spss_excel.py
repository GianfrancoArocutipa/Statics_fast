# -*- coding: utf-8 -*-
"""
Generador de la Guía Paso a Paso en Excel para la Segunda Práctica Calificada
Personalizado exactamente con los 5 problemas, variables, botones y opciones de SPSS.
"""

import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def create_exam_excel_guide(output_path=os.path.join("prueba", "Guia_Pasos_SPSS_Segunda_Practica.xlsx")):
    wb = openpyxl.Workbook()
    
    # Fuentes y estilos
    font_title = Font(name="Calibri", size=13, bold=True, color="103B60")
    font_subtitle = Font(name="Calibri", size=10, italic=True, color="4A5568")
    font_header = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    font_body = Font(name="Calibri", size=10, color="1A202C")
    font_bold = Font(name="Calibri", size=10, bold=True, color="1A202C")
    
    fill_header = PatternFill(start_color="103B60", end_color="103B60", fill_type="solid")
    fill_zebra = PatternFill(start_color="F7FAFC", end_color="F7FAFC", fill_type="solid")
    fill_white = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    fill_btn = PatternFill(start_color="EBF4FF", end_color="EBF4FF", fill_type="solid")

    thin_border_side = Side(border_style="thin", color="CBD5E0")
    thin_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)

    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    align_header = Alignment(horizontal="center", vertical="center", wrap_text=True)

    headers = [
        "Paso Nº",
        "Fase / Objetivo",
        "Ruta en Menú SPSS",
        "Ventana / Cuadro de Diálogo",
        "Botón o Pestaña Específica",
        "Acción Específica (Casillas, Parámetros y Campos)",
        "Resultado Esperado / Observaciones"
    ]

    col_widths = [8, 22, 28, 26, 25, 45, 30]

    # ==========================================
    # HOJA 1: RESUMEN DE LA PRÁCTICA
    # ==========================================
    ws_resumen = wb.active
    ws_resumen.title = "Resumen_Practica"
    ws_resumen.views.sheetView[0].showGridLines = True

    ws_resumen.merge_cells("A1:G1")
    ws_resumen["A1"] = "SEGUNDA PRÁCTICA CALIFICADA - RESUMEN DE EJERCICIOS Y PRUEBAS SPSS"
    ws_resumen["A1"].font = font_title
    ws_resumen["A1"].alignment = Alignment(horizontal="left", vertical="center")

    ws_resumen.merge_cells("A2:G2")
    ws_resumen["A2"] = "Universidad Privada de Tacna | Escuela de Ingeniería de Sistemas | Dr. Luis A. Fernández Vizcarra"
    ws_resumen["A2"].font = font_subtitle
    ws_resumen["A2"].alignment = Alignment(horizontal="left", vertical="center")

    res_headers = ["Ejercicio", "Contexto y Variables", "Nivel Sig. (α)", "Prueba Estadística Aplicada", "Estadístico y p-valor", "Decisión Estadística", "Conclusión Breve"]
    for c_idx, h_text in enumerate(res_headers, 1):
        cell = ws_resumen.cell(row=4, column=c_idx, value=h_text)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_header
        cell.border = thin_border

    res_data = [
        (
            "Ejercicio 1 (4 ptos)",
            "Incidentes reportados en app web (4 categorías: Autenticación, Conexión, Rendimiento, Interfaz). N=120.",
            "3% (0.03)",
            "Prueba no paramétrica Chi-cuadrado de bondad de ajuste",
            "Chi2 = 9.267, gl = 3, p = 0.026",
            "Rechazar H0 (p = 0.026 < 0.03)",
            "A un margen del 3% se concluye que los incidentes reportados no se distribuyen uniformemente entre las 4 categorías."
        ),
        (
            "Ejercicio 2 (4 ptos)",
            "Tiempos de ejecución (ms) de dos algoritmos de búsqueda: Algoritmo A (n=10) vs Algoritmo B (n=11).",
            "1% (0.01)",
            "Shapiro-Wilk (Normales: p=0.957 y 0.997) -> T de Student para muestras independientes",
            "t = 5.556, gl = 19, p = 0.000023 (Mann-Whitney U=3.0, p=0.000)",
            "Rechazar H0 (p < 0.01)",
            "A un margen del 1% se concluye que los tiempos de ejecución difieren significativamente entre ambos algoritmos."
        ),
        (
            "Ejercicio 3 (4 ptos)",
            "Secuencia de 30 intentos consecutivos de acceso clasificados como Éxito (E, 16) o Fallo (F, 14).",
            "5% (0.05)",
            "Prueba no paramétrica de Rachas (Runs test de aleatoriedad)",
            "Rachas = 21, Z = 1.891, p = 0.059",
            "No rechazar H0 (p = 0.059 > 0.05)",
            "A un margen del 5% se concluye que la secuencia de intentos de acceso al sistema presenta un comportamiento aleatorio."
        ),
        (
            "Ejercicio 4 (4 ptos)",
            "Transacciones electrónicas (n=50, 40 exitosas y 10 con error). Estándar histórico: 90% (0.90).",
            "5% (0.05)",
            "Prueba no paramétrica Binomial (unilateral inferior)",
            "Prop. obs = 0.80, Prop. prueba = 0.90, p = 0.025",
            "Rechazar H0 (p = 0.025 < 0.05)",
            "A un margen del 5% se concluye que la proporción de transacciones exitosas es inferior al estándar del 90%."
        ),
        (
            "Ejercicio 5 (4 ptos)",
            "Tiempos de respuesta (ms) de 3 servicios cloud: Servicio A (n=11), Servicio B (n=13) y Servicio C (n=15).",
            "3% (0.03)",
            "Shapiro-Wilk (Normales) + Levene (p=0.625) -> ANOVA de un factor + Post Hoc Tukey",
            "F = 316.629, gl = (2, 36), p = 0.000 (Kruskal H=33.60, p=0.000)",
            "Rechazar H0 (p < 0.03)",
            "A un margen del 3% se concluye que, por lo menos los tiempos de respuesta de 2 servicios de almacenamiento en la nube son diferentes."
        )
    ]

    for r_idx, row_items in enumerate(res_data, 5):
        fill_curr = fill_zebra if r_idx % 2 == 1 else fill_white
        ws_resumen.row_dimensions[r_idx].height = 48
        for c_idx, val in enumerate(row_items, 1):
            cell = ws_resumen.cell(row=r_idx, column=c_idx, value=val)
            cell.font = font_bold if c_idx in [1, 4] else font_body
            cell.fill = fill_curr
            cell.alignment = align_center if c_idx in [1, 3] else align_left
            cell.border = thin_border

    res_widths = [18, 30, 16, 28, 26, 22, 35]
    for i, w in enumerate(res_widths, 1):
        ws_resumen.column_dimensions[get_column_letter(i)].width = w

    # ==========================================
    # HOJAS DE DETALLE DE CADA EJERCICIO
    # ==========================================
    exercise_sheets = [
        {
            "code": "Ejer1_Chi_Cuadrado",
            "title": "EJERCICIO 1: CHI-CUADRADO DE BONDAD DE AJUSTE (DISTRIBUCIÓN DE INCIDENTES)",
            "desc": "Evaluar si 120 incidentes reportados se distribuyen uniformemente entre 4 categorías (α = 3%).",
            "steps": [
                ("1", "Definir Variables", "Vista de variables", "Fila 1 y Fila 2", "Teclado", "Fila 1: Nombre = 'Categoria', Tipo = Numérico, Decimales = 0, Medida = Nominal. Fila 2: Nombre = 'Frecuencia', Tipo = Numérico, Decimales = 0, Medida = Escala.", "Variables creadas en SPSS."),
                ("2", "Etiquetar Categorías", "Vista de variables (Fila 1)", "Columna 'Valores'", "Botón [...]", "Clic en [...]. Añadir: 1 = 'Errores de autenticación', 2 = 'Fallas de conexión', 3 = 'Problemas de rendimiento', 4 = 'Errores de interfaz' -> Clic en [Aceptar].", "Categorías codificadas."),
                ("3", "Ingresar Datos", "Vista de datos", "Filas 1 a 4", "Teclado", "Fila 1: Categoria=1, Frecuencia=42. Fila 2: Categoria=2, Frecuencia=25. Fila 3: Categoria=3, Frecuencia=33. Fila 4: Categoria=4, Frecuencia=20.", "Matriz de frecuencias cargada."),
                ("4", "Ponderar Casos", "Barra Superior > Datos", "Menú Desplegable", "Ponderar casos...", "Ir a 'Datos' > 'Ponderar casos...'. Marcar 'Ponderar casos mediante'. Mover 'Frecuencia' con [ ► ] hacia 'Variable de frecuencia' -> Clic en [Aceptar].", "SPSS activa la ponderación: 'Ponderación activada'."),
                ("5", "Abrir Prueba Chi-cuadrado", "Barra Superior > Analizar", "Menú Desplegable", "Pruebas no paramétricas > Cuadros de diálogo antiguos > Chi-cuadrado...", "Ir a 'Analizar' > 'Pruebas no paramétricas' > 'Cuadros de diálogo antiguos' > 'Chi-cuadrado...'.", "Se abre el cuadro 'Prueba de chi-cuadrado'."),
                ("6", "Asignar Variable a Contrastar", "Cuadro: Prueba de chi-cuadrado", "Panel izquierdo", "Botón Flecha [ ► ]", "Seleccionar 'Categoria' y hacer clic en la flecha [ ► ] apuntando a 'Lista contrastar variables'.", "'Categoria' queda en la lista de prueba."),
                ("7", "Verificar Valores Esperados", "Cuadro: Prueba de chi-cuadrado", "Sección 'Valores esperados'", "Botón de opción (Radio button)", "Verificar que esté marcada la opción: '(o) Todas las categorías iguales'.", "Distribución uniforme esperada (30 incidentes por categoría)."),
                ("8", "Ejecutar Prueba", "Cuadro: Prueba de chi-cuadrado", "Panel derecho", "Botón [Aceptar]", "Hacer clic en el botón [Aceptar].", "Se generan las tablas: 'Categorías de incidentes reportados' y 'Estadísticos de prueba' (Chi-cuadrado = 9.267, gl = 3, Sig. asintót. = 0.026). Como p=0.026 < 0.03, se rechaza H0.")
            ]
        },
        {
            "code": "Ejer2_Algoritmos_A_B",
            "title": "EJERCICIO 2: COMPARACIÓN DE 2 ALGORITMOS INDEPENDIENTES (TIEMPOS EN MS)",
            "desc": "Comparar tiempos de ejecución entre Algoritmo A (n=10) y Algoritmo B (n=11) al nivel α = 1%.",
            "steps": [
                ("1", "Definir Variables", "Vista de variables", "Fila 1 y Fila 2", "Teclado", "Fila 1: Nombre = 'Algoritmo', Decimales = 0, Medida = Nominal. Fila 2: Nombre = 'Tiempo_ms', Decimales = 0, Medida = Escala.", "Estructura de datos lista."),
                ("2", "Etiquetar Valores", "Vista de variables (Fila 1)", "Columna 'Valores'", "Botón [...]", "Clic en [...]. Añadir: 1 = 'Algoritmo A', 2 = 'Algoritmo B' -> Clic en [Aceptar].", "Etiquetas de grupo asignadas."),
                ("3", "Ingresar Datos", "Vista de datos", "Filas 1 a 21", "Teclado / Pegar", "Filas 1-10: Algoritmo=1 y Tiempos de A (120, 135, 128, 142, 130, 125, 138, 132, 145, 127). Filas 11-21: Algoritmo=2 y Tiempos de B (110, 118, 115, 122, 108, 120, 112, 116, 125, 119, 114).", "Total de 21 observaciones cargadas."),
                ("4", "Descriptivos (Informe)", "Analizar > Comparar medias > Medias...", "Cuadro: Medias", "Botón [Opciones...]", "Mover 'Tiempo_ms' a Dependientes y 'Algoritmo' a Independientes. En [Opciones...] asegurar Media, N, Desv. típica -> Clic en [Continuar] -> Clic en [Aceptar].", "Genera tabla 'Informe' (Media A = 132.20 ms, Media B = 116.27 ms)."),
                ("5", "Prueba de Normalidad", "Analizar > Estadísticos descriptivos > Explorar...", "Cuadro: Explorar", "Botón [Gráficos...]", "Mover 'Tiempo_ms' a Dependientes y 'Algoritmo' a Factores. Clic en [Gráficos...], MARCAR '[X] Gráficos de normalidad con pruebas' -> Clic en [Continuar] -> Clic en [Aceptar].", "Genera tabla 'Pruebas de normalidad': Shapiro-Wilk A (p=0.957) y B (p=0.997). Ambos p > 0.01 -> Distribución normal."),
                ("6", "Ejecutar Prueba T", "Analizar > Comparar medias > Prueba T para muestras independientes...", "Cuadro: Prueba T", "Botón [Definir grupos...]", "Mover 'Tiempo_ms' a 'Variables de prueba' y 'Algoritmo' a 'Variable de agrupación'. Clic en [Definir grupos...], ingresar Grupo 1: 1 y Grupo 2: 2 -> [Continuar] -> Clic en [Aceptar].", "Genera tabla 'Prueba de muestras independientes': Levene (p=0.189), t = 5.556, gl = 19, Sig. (bilateral) = 0.000. Como p < 0.01, se rechaza H0."),
                ("7", "Prueba Alternativa U Mann-Whitney", "Analizar > Pruebas no paramétricas > Cuadros antiguos > 2 muestras independientes...", "Cuadro: 2 muestras indep.", "Botón [Definir grupos...]", "Mover 'Tiempo_ms' a contraste y 'Algoritmo' a agrupación. [Definir grupos...] (1 y 2). Marcar '[X] U de Mann-Whitney' -> Clic en [Aceptar].", "Genera 'Estadísticos de prueba': U = 3.000, W = 69.000, Z = -3.535, Sig. = 0.000 (rechaza H0).")
            ]
        },
        {
            "code": "Ejer3_Prueba_Rachas",
            "title": "EJERCICIO 3: PRUEBA DE RACHAS (ALEATORIEDAD DE INTENTOS DE ACCESO)",
            "desc": "Evaluar la aleatoriedad de 30 intentos de acceso (16 Éxitos y 14 Fallos) al nivel α = 5%.",
            "steps": [
                ("1", "Definir Variable", "Vista de variables", "Fila 1", "Teclado", "Nombre = 'Acceso', Tipo = Numérico, Decimales = 0, Medida = Nominal.", "Variable creada."),
                ("2", "Etiquetar Valores", "Vista de variables (Fila 1)", "Columna 'Valores'", "Botón [...]", "Clic en [...]. Añadir: 1 = 'Acceso exitoso (E)', 0 = 'Acceso fallido (F)' -> Clic en [Aceptar].", "Valores dicotómicos codificados."),
                ("3", "Ingresar Secuencia", "Vista de datos", "Filas 1 a 30", "Teclado", "Digitar en orden la secuencia de 30 datos: 1, 1, 0, 1, 0, 0, 1, 1, 0, 1, 0, 0, 1, 0, 1, 1, 0, 0, 1, 0, 1, 1, 0, 1, 0, 0, 1, 1, 0, 1.", "Secuencia de 30 observaciones ingresada."),
                ("4", "Abrir Prueba de Rachas", "Barra Superior > Analizar", "Menú Desplegable", "Pruebas no paramétricas > Cuadros de diálogo antiguos > Rachas...", "Ir a 'Analizar' > 'Pruebas no paramétricas' > 'Cuadros de diálogo antiguos' > 'Rachas...'.", "Se abre el cuadro de diálogo 'Prueba de rachas'."),
                ("5", "Asignar Variable", "Cuadro: Prueba de rachas", "Panel izquierdo", "Botón Flecha [ ► ]", "Seleccionar 'Acceso' y hacer clic en la flecha [ ► ] hacia 'Lista contrastar variables'.", "'Acceso' queda en la lista de prueba."),
                ("6", "Configurar Punto de Corte", "Cuadro: Prueba de rachas", "Sección 'Punto de corte'", "Casillas de verificación", "Verificar que esté marcada la opción: '[X] Mediana' (o marcar '[X] Personalizado' y colocar '0.5').", "Define la separación entre éxitos y fallos."),
                ("7", "Ejecutar Prueba", "Cuadro: Prueba de rachas", "Panel derecho", "Botón [Aceptar]", "Hacer clic en el botón [Aceptar].", "Se genera la tabla 'Prueba de rachas' (Total casos = 30, Rachas = 21, Z = 1.891, Sig. asintót. bilateral = 0.059). Como p=0.059 > 0.05, NO se rechaza H0 (comportamiento aleatorio).")
            ]
        },
        {
            "code": "Ejer4_Prueba_Binomial",
            "title": "EJERCICIO 4: PRUEBA BINOMIAL (PROPORCIÓN DE TRANSACCIONES EXITOSAS)",
            "desc": "Evaluar si la proporción de transacciones exitosas es inferior al estándar del 90% (α = 5%). Muestra n=50.",
            "steps": [
                ("1", "Definir Variables", "Vista de variables", "Fila 1 y Fila 2", "Teclado", "Fila 1: Nombre = 'Estado', Tipo = Numérico, Decimales = 0, Medida = Nominal. Fila 2: Nombre = 'Cantidad', Tipo = Numérico, Decimales = 0, Medida = Escala.", "Variables creadas."),
                ("2", "Etiquetar Valores", "Vista de variables (Fila 1)", "Columna 'Valores'", "Botón [...]", "Clic en [...]. Añadir: 1 = 'Transacción exitosa', 0 = 'Transacción con error' -> Clic en [Aceptar].", "Estados codificados."),
                ("3", "Ingresar Datos Resumidos", "Vista de datos", "Filas 1 y 2", "Teclado", "Fila 1: Estado = 1, Cantidad = 40. Fila 2: Estado = 0, Cantidad = 10.", "Datos ingresados."),
                ("4", "Ponderar Casos", "Barra Superior > Datos", "Menú Desplegable", "Ponderar casos...", "Ir a 'Datos' > 'Ponderar casos...'. Marcar 'Ponderar casos mediante'. Mover 'Cantidad' con [ ► ] hacia 'Variable de frecuencia' -> Clic en [Aceptar].", "Ponderación activada para n=50 casos."),
                ("5", "Abrir Prueba Binomial", "Barra Superior > Analizar", "Menú Desplegable", "Pruebas no paramétricas > Cuadros de diálogo antiguos > Binomial...", "Ir a 'Analizar' > 'Pruebas no paramétricas' > 'Cuadros de diálogo antiguos' > 'Binomial...'.", "Se abre el cuadro de diálogo 'Prueba binomial'."),
                ("6", "Asignar Variable a Contrastar", "Cuadro: Prueba binomial", "Panel izquierdo", "Botón Flecha [ ► ]", "Seleccionar 'Estado' y mover con [ ► ] hacia 'Lista contrastar variables'.", "'Estado' queda asignado."),
                ("7", "Ingresar Proporción de Prueba", "Cuadro: Prueba binomial", "Campo 'Proporción de prueba'", "Teclado", "Borrar 0.50 y escribir obligatoriamente: '0.90' (el estándar del 90%).", "Proporción de prueba establecida en 0.90."),
                ("8", "Ejecutar Prueba", "Cuadro: Prueba binomial", "Panel derecho", "Botón [Aceptar]", "Hacer clic en el botón [Aceptar].", "Se genera la tabla 'Prueba binomial': N=40 exitosas (80%), N=10 errores (20%), Prop. prueba = 0.90, Sig. exacta (unilateral) = 0.025. Como p=0.025 < 0.05, se rechaza H0.")
            ]
        },
        {
            "code": "Ejer5_Cloud_ANOVA",
            "title": "EJERCICIO 5: ANOVA Y POST HOC (3 SERVICIOS DE ALMACENAMIENTO EN LA NUBE)",
            "desc": "Comparar tiempos de respuesta (ms) entre Servicio A, Servicio B y Servicio C al nivel α = 3%.",
            "steps": [
                ("1", "Definir Variables", "Vista de variables", "Fila 1 y Fila 2", "Teclado", "Fila 1: Nombre = 'Servicio', Tipo = Numérico, Decimales = 0, Medida = Nominal. Fila 2: Nombre = 'Tiempo_ms', Tipo = Numérico, Decimales = 0, Medida = Escala.", "Variables listas."),
                ("2", "Etiquetar Servicios", "Vista de variables (Fila 1)", "Columna 'Valores'", "Botón [...]", "Clic en [...]. Añadir: 1 = 'Servicio A', 2 = 'Servicio B', 3 = 'Servicio C' -> Clic en [Aceptar].", "Servicios codificados: 1, 2 y 3."),
                ("3", "Ingresar Datos", "Vista de datos", "Filas 1 a 39", "Teclado / Pegar", "Filas 1-11: Servicio=1 y Tiempos A. Filas 12-24: Servicio=2 y Tiempos B. Filas 25-39: Servicio=3 y Tiempos C.", "Total 39 registros cargados."),
                ("4", "Descriptivos (Informe)", "Analizar > Comparar medias > Medias...", "Cuadro: Medias", "Botón [Opciones...]", "Mover 'Tiempo_ms' a Dependientes y 'Servicio' a Independientes. En [Opciones...] asegurar Media, N, Desv. típica -> Clic en [Continuar] -> Clic en [Aceptar].", "Genera tabla 'Informe' (Media A = 100.18, Media B = 121.08, Media C = 141.00 ms)."),
                ("5", "Prueba de Normalidad", "Analizar > Estadísticos descriptivos > Explorar...", "Cuadro: Explorar", "Botón [Gráficos...]", "Mover 'Tiempo_ms' a Dependientes y 'Servicio' a Factores. En [Gráficos...] marcar '[X] Gráficos de normalidad con pruebas' -> Clic en [Continuar] -> Clic en [Aceptar].", "Genera tabla 'Pruebas de normalidad': Shapiro-Wilk A (p=0.885), B (p=0.954), C (p=0.755). Todos p > 0.03 -> Distribución normal."),
                ("6", "Configurar Opciones ANOVA", "Analizar > Comparar medias > ANOVA de un factor...", "Cuadro: ANOVA de un factor", "Botón [Opciones...]", "Mover 'Tiempo_ms' a 'Lista de dependientes' y 'Servicio' a 'Factor'. Clic en botón [Opciones...], marcar '[X] Descriptivos' y '[X] Prueba de homogeneidad de las varianzas' -> Clic en [Continuar].", "Se activa la prueba de homogeneidad de varianzas de Levene."),
                ("7", "Configurar Post Hoc Tukey", "Cuadro: ANOVA de un factor", "Panel derecho", "Botón [Post Hoc...]", "Hacer clic en el botón [Post Hoc...]. En la sección 'Se asumen varianzas iguales', MARCAR la casilla: '[X] Tukey'. Verificar 'Nivel de significación: 0.05' -> Clic en [Continuar].", "Prueba Post Hoc de Tukey configurada."),
                ("8", "Ejecutar ANOVA", "Cuadro: ANOVA de un factor", "Panel derecho", "Botón [Aceptar]", "Hacer clic en el botón [Aceptar].", "Se generan las tablas: 'Prueba de homogeneidad de varianzas' (Levene p = 0.625), 'ANOVA' (F = 316.629, Sig. = 0.000) y 'Comparaciones múltiples' (Tukey: A vs B p=0.000, A vs C p=0.000, B vs C p=0.000). Se rechaza H0."),
                ("9", "Prueba Kruskal-Wallis (Alternativa)", "Analizar > Pruebas no paramétricas > Cuadros antiguos > K muestras independientes...", "Cuadro: K muestras indep.", "Botón [Definir rango...]", "Mover 'Tiempo_ms' a prueba y 'Servicio' a agrupación. [Definir rango...] Mínimo: 1, Máximo: 3 -> [Continuar]. Marcar '[X] Kruskal-Wallis H' -> Clic en [Aceptar].", "Genera tabla 'Estadísticos de prueba': H = 33.600, gl = 2, Sig. asintót. = 0.000 (rechaza H0).")
            ]
        },
        {
            "code": "Guardar_SAV_y_SPV",
            "title": "GUARDADO DE ARCHIVOS DE DATOS (.SAV) Y RESULTADOS (.SPV)",
            "desc": "Procedimiento exacto para guardar cada ejercicio según los estándares solicitados por el docente.",
            "steps": [
                ("1", "Guardar Datos Ejercicio 1", "Ventana Editor de Datos SPSS", "Archivo > Guardar como...", "Botón [Guardar]", "Nombre de archivo: 'Ejercicio01_GF.sav'. Guardar como tipo: 'SPSS Statistics (*.sav)'. Clic en [Guardar].", "Archivo .sav de datos creado."),
                ("2", "Guardar Visor Ejercicio 1", "Ventana Visor de Resultados SPSS", "Archivo > Guardar como...", "Botón [Guardar]", "Nombre de archivo: 'Ejercicio01.spv'. Guardar como tipo: 'Visor de SPSS Statistics (*.spv)'. Clic en [Guardar].", "Archivo .spv de resultados creado."),
                ("3", "Guardar Datos Ejercicio 2", "Ventana Editor de Datos SPSS", "Archivo > Guardar como...", "Botón [Guardar]", "Nombre de archivo: 'Ejercicio02_GF.sav'. Guardar como tipo: 'SPSS Statistics (*.sav)'. Clic en [Guardar].", "Archivo .sav de datos creado."),
                ("4", "Guardar Visor Ejercicio 2", "Ventana Visor de Resultados SPSS", "Archivo > Guardar como...", "Botón [Guardar]", "Nombre de archivo: 'Ejercicio02.spv'. Guardar como tipo: 'Visor de SPSS Statistics (*.spv)'. Clic en [Guardar].", "Archivo .spv de resultados creado."),
                ("5", "Guardar Datos Ejercicio 3", "Ventana Editor de Datos SPSS", "Archivo > Guardar como...", "Botón [Guardar]", "Nombre de archivo: 'Ejercicio03_GF.sav'. Guardar como tipo: 'SPSS Statistics (*.sav)'. Clic en [Guardar].", "Archivo .sav de datos creado."),
                ("6", "Guardar Visor Ejercicio 3", "Ventana Visor de Resultados SPSS", "Archivo > Guardar como...", "Botón [Guardar]", "Nombre de archivo: 'Ejercicio03.spv'. Guardar como tipo: 'Visor de SPSS Statistics (*.spv)'. Clic en [Guardar].", "Archivo .spv de resultados creado."),
                ("7", "Guardar Datos Ejercicio 4", "Ventana Editor de Datos SPSS", "Archivo > Guardar como...", "Botón [Guardar]", "Nombre de archivo: 'Ejercicio04_GF.sav'. Guardar como tipo: 'SPSS Statistics (*.sav)'. Clic en [Guardar].", "Archivo .sav de datos creado."),
                ("8", "Guardar Visor Ejercicio 4", "Ventana Visor de Resultados SPSS", "Archivo > Guardar como...", "Botón [Guardar]", "Nombre de archivo: 'Ejercicio04.spv'. Guardar como tipo: 'Visor de SPSS Statistics (*.spv)'. Clic en [Guardar].", "Archivo .spv de resultados creado."),
                ("9", "Guardar Datos Ejercicio 5", "Ventana Editor de Datos SPSS", "Archivo > Guardar como...", "Botón [Guardar]", "Nombre de archivo: 'Ejercicio05_GF.sav'. Guardar como tipo: 'SPSS Statistics (*.sav)'. Clic en [Guardar].", "Archivo .sav de datos creado."),
                ("10", "Guardar Visor Ejercicio 5", "Ventana Visor de Resultados SPSS", "Archivo > Guardar como...", "Botón [Guardar]", "Nombre de archivo: 'Ejercicio05.spv'. Guardar como tipo: 'Visor de SPSS Statistics (*.spv)'. Clic en [Guardar].", "Archivo .spv de resultados creado.")
            ]
        }
    ]

    for s_info in exercise_sheets:
        ws = wb.create_sheet(title=s_info["code"])
        ws.views.sheetView[0].showGridLines = True

        ws.merge_cells("A1:G1")
        ws["A1"] = s_info["title"]
        ws["A1"].font = font_title
        ws["A1"].alignment = Alignment(horizontal="left", vertical="center")

        ws.merge_cells("A2:G2")
        ws["A2"] = s_info["desc"]
        ws["A2"].font = font_subtitle
        ws["A2"].alignment = Alignment(horizontal="left", vertical="center")

        for c_idx, h_text in enumerate(headers, 1):
            cell = ws.cell(row=4, column=c_idx, value=h_text)
            cell.font = font_header
            cell.fill = fill_header
            cell.alignment = align_header
            cell.border = thin_border
        
        ws.row_dimensions[4].height = 28

        for r_idx, step_data in enumerate(s_info["steps"], 5):
            fill_curr = fill_zebra if r_idx % 2 == 1 else fill_white
            ws.row_dimensions[r_idx].height = 42

            for c_idx, val in enumerate(step_data, 1):
                cell = ws.cell(row=r_idx, column=c_idx, value=val)
                cell.fill = fill_curr
                cell.border = thin_border

                if c_idx == 1:
                    cell.font = font_bold
                    cell.alignment = align_center
                elif c_idx == 5:
                    cell.font = font_bold
                    cell.alignment = align_left
                    cell.fill = fill_btn
                elif c_idx in [2, 3, 4]:
                    cell.font = font_bold if c_idx == 2 else font_body
                    cell.alignment = align_left
                else:
                    cell.font = font_body
                    cell.alignment = align_left

        for i, w in enumerate(col_widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

    wb.save(output_path)
    print(f"Excel generado exitosamente en: {output_path}")

if __name__ == "__main__":
    create_exam_excel_guide()
