import os
import docx
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = docx.Document()

# Set A4 margins
for section in doc.sections:
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    section.header_distance = Inches(0.4)
    
    # Configure Header on every page
    header = section.header
    hp = header.paragraphs[0]
    hp.paragraph_format.space_before = Pt(0)
    hp.paragraph_format.space_after = Pt(0)
    hp.paragraph_format.line_spacing = 1.05
    
    r1 = hp.add_run("Universidad Privada de Tacna\n")
    r1.bold = True
    r1.font.name = "Calibri"
    r1.font.size = Pt(10)
    
    r2 = hp.add_run("Facultad de Ingeniería\n")
    r2.bold = True
    r2.font.name = "Calibri"
    r2.font.size = Pt(9.5)
    
    r3 = hp.add_run("Escuela Profesional de Ingeniería de Sistemas")
    r3.bold = True
    r3.font.name = "Calibri"
    r3.font.size = Pt(9)

def add_title_block(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(12)
    p.paragraph_format.line_spacing = 1.15
    r1 = p.add_run("Ejercicios de prueba\n")
    r1.bold = True
    r1.font.name = "Calibri"
    r1.font.size = Pt(13)
    r2 = p.add_run("Análisis de la varianza de Kruskal-Wallis")
    r2.bold = True
    r2.font.name = "Calibri"
    r2.font.size = Pt(12)

def add_p(doc, text="", bold_prefix="", space_before=1, space_after=3, align=WD_ALIGN_PARAGRAPH.LEFT):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.12
    if bold_prefix:
        r_b = p.add_run(bold_prefix)
        r_b.bold = True
        r_b.font.name = "Calibri"
        r_b.font.size = Pt(10.5)
    if text:
        r_t = p.add_run(text)
        r_t.font.name = "Calibri"
        r_t.font.size = Pt(10.5)
    return p

def add_img(doc, img_path, width=Inches(3.2), space_before=2, space_after=4):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    r = p.add_run()
    r.add_picture(img_path, width=width)
    return p

# --- BUILD DOCUMENT ---
add_title_block(doc)

# ==========================================
# EJERCICIO 1
# ==========================================
add_p(doc, "1. Se desea comparar el tiempo de respuesta (en milisegundos) de tres sistemas operativos (X, Y y Z) bajo la misma carga de trabajo. Se han realizado pruebas para cada sistema y se obtuvieron los siguientes tiempos de respuesta:\nSistema X: 120  115  130  125  110  118\nSistema Y: 150  145  148  140  142  144\nSistema Z: 170  168  130  160  162  169  165\n¿Existe una diferencia significativa en el tiempo de respuesta de los sistemas operativos? Contrastar la hipótesis a un margen de error del 5%.", space_after=4)

add_p(doc, bold_prefix="Solución:")
add_p(doc, bold_prefix="Paso 1: Formación de hipótesis")
add_p(doc, "Las Promedios de los grupos son iguales", bold_prefix="H0: ")
add_p(doc, "x̄i = x̄y")
add_p(doc, "Por lo menos los Promedios de 2 grupos son diferentes", bold_prefix="H1: ")
add_p(doc, "x̄i ≠ x̄y")

add_p(doc, bold_prefix="Paso 2: Determinar el nivel de significancia")
add_p(doc, "N.S = 5% (0.05)")

add_p(doc, bold_prefix="Paso 3: Hallar estadísticos descriptivos")
if os.path.exists("extracted_kw_pdf_images/page_1_Image25.png"):
    add_img(doc, "extracted_kw_pdf_images/page_1_Image25.png", width=Inches(3.1))

add_p(doc, bold_prefix="Paso 4: Prueba de Distribución normal")
add_p(doc, "Los datos de los grupos X, Y y Z presentan distribución normal", bold_prefix="H0: ")
add_p(doc, "Por lo menos los datos de un grupo X, Y o Z no presentan distribución normal", bold_prefix="H1: ")
if os.path.exists("extracted_kw_pdf_images/page_2_Image28.png"):
    add_img(doc, "extracted_kw_pdf_images/page_2_Image28.png", width=Inches(3.5))

add_p(doc, "Dado que los p-valores de los grupos X (0.991) e Y (0.938) son mayores que el nivel de significancia, pero el grupo Z (0.003) es menor al nivel de significancia por lo que es probable que se rechace la hipótesis nula")
add_p(doc, "A un margen de error de 5 % se concluye que , los datos de los grupos X, Y y Z, no presentan Distribución normal", bold_prefix="Conclusión: ")

add_p(doc, bold_prefix="Paso 5: Elegir la prueba estadística")
add_p(doc, "Dado que los datos del grupo Z no presentan distribución normal  se aplicara la prueba Kruskal-Wallis (Prueba no paramétrica)")
if os.path.exists("extracted_kw_pdf_images/page_2_Image29.png"):
    add_img(doc, "extracted_kw_pdf_images/page_2_Image29.png", width=Inches(2.8))
if os.path.exists("extracted_kw_pdf_images/page_2_Image30.png"):
    add_img(doc, "extracted_kw_pdf_images/page_2_Image30.png", width=Inches(2.4))

add_p(doc, bold_prefix="Paso 6: Contrastación de Hipótesis")
add_p(doc, "Dado que el p -valor de Kruskal -Wallis (0.001) es menor que el nivel de significancia (0.05), se rechaza la hipótesis nula.")

add_p(doc, bold_prefix="Paso 7: Conclusión")
add_p(doc, "A un margen de error del 5% se concluye que, por lo menos los tiempos de respuesta de 2 sistemas operativos son diferentes.", space_after=12)

# ==========================================
# EJERCICIO 2
# ==========================================
add_p(doc, "2. Se desea comparar el rendimiento de tres algoritmos de clasificación (A, B y C) en términos de su precisión sobre un conjunto de datos. Se han realizado pruebas para cada algoritmo y se obtuvieron los siguientes porcentajes de precisión:\nAlgoritmo A: 85  87  84  86  88\nAlgoritmo B: 78  79  81  80  77  85  87  99\nAlgoritmo C: 92  91  93  94  90  87\n¿Existe una diferencia significativa en el rendimiento de los algoritmos? Contrastar la hipótesis a un margen de error del 3%.", space_after=4)

add_p(doc, bold_prefix="Solución:")
add_p(doc, bold_prefix="Paso 1: Formación de hipótesis")
add_p(doc, "Las Promedios de los grupos son iguales", bold_prefix="H0: ")
add_p(doc, "x̄i = x̄")
add_p(doc, "Por lo menos los Promedios de 2 grupos son diferentes", bold_prefix="H1: ")

add_p(doc, bold_prefix="Paso 2: Determinar el nivel de significancia")
add_p(doc, "N.S = 3% (0.03)")

add_p(doc, bold_prefix="Paso 3: Hallar estadísticos descriptivos")
if os.path.exists("extracted_kw_pdf_images/page_3_Image33.png"):
    add_img(doc, "extracted_kw_pdf_images/page_3_Image33.png", width=Inches(3.1))

add_p(doc, bold_prefix="Paso 4: Prueba de Distribución normal")
add_p(doc, "Los datos de los grupos A, B y C presentan distribución normal", bold_prefix="H0: ")
add_p(doc, "Por lo menos los datos de un grupo A, B o C no presentan distribución normal", bold_prefix="H1: ")
if os.path.exists("extracted_kw_pdf_images/page_4_Image36.png"):
    add_img(doc, "extracted_kw_pdf_images/page_4_Image36.png", width=Inches(3.5))

add_p(doc, "Dado que los p -valores de los grupos A (0.967), B (0.042) y C (0.794) son mayores que el nivel de significancia, es probable que no se rechace la hipótesis nula")
add_p(doc, "A un margen de error de 3 % se concluye que, los datos de los grupos A, B y C, presentan Distribución normal", bold_prefix="Conclusión: ")

add_p(doc, bold_prefix="Paso 5.A: Prueba de homogeneidad de varianzas")
add_p(doc, "Las varianzas de los grupos A, B y C son iguales", bold_prefix="H0: ")
add_p(doc, "Las varianzas de los grupos A, B o C son diferentes", bold_prefix="H1: ")
if os.path.exists("extracted_kw_pdf_images/page_4_Image37.png"):
    add_img(doc, "extracted_kw_pdf_images/page_4_Image37.png", width=Inches(3.8))

add_p(doc, "A un margen de error del 3% se concluye que, Las varianzas de los grupos A, B y C son iguales.", bold_prefix="Conclusión: ")

add_p(doc, bold_prefix="Paso 6: Elegir la prueba estadística")
add_p(doc, "Dado que los datos de los 3 grupos presentan distribución normal se aplicará la prueba de ANOVA (Prueba paramétrica)")

add_p(doc, bold_prefix="Paso 7: Contrastación de Hipótesis")
if os.path.exists("extracted_kw_pdf_images/page_4_Image38.png"):
    add_img(doc, "extracted_kw_pdf_images/page_4_Image38.png", width=Inches(3.8))

add_p(doc, "Dado que el p-valor de ANOVA (0.033), es mayor que el nivel de significancia (0.03) es probable que no se rechace la hipótesis nula")

add_p(doc, bold_prefix="Paso 8: Conclusión")
add_p(doc, "A un margen de error del 3 % se concluye que, Las Promedios de los grupos son iguales.\n\nSin embargo, de acuerdo a la tabla de comparaciones múltiples se a podido identificar que las puntuaciones promedias de los grupos B y C son diferentes (p-valor = 0.026).")
if os.path.exists("extracted_kw_pdf_images/page_5_Image41.png"):
    add_img(doc, "extracted_kw_pdf_images/page_5_Image41.png", width=Inches(4.4), space_after=12)

# ==========================================
# EJERCICIO 3 (RESUELTO)
# ==========================================
add_p(doc, "3. Se desea comparar el consumo de memoria (en MB) de tres aplicaciones (1, 2 y 3) durante su ejecución. Se han realizado pruebas para cada aplicación y se obtuvieron los siguientes consumos de memoria:\nAplicación 1: 350  360  355  340  315  285  245\nAplicación 2: 410  420  415  405\nAplicación 3: 300  310  305  295  283  225\n¿Existe una diferencia significativa en el consumo de memoria de las aplicaciones? Contrastar la hipótesis a un margen de error del 5%.", space_after=4)

add_p(doc, bold_prefix="Solución:")
add_p(doc, bold_prefix="Paso 1: Formación de hipótesis")
add_p(doc, "Las Promedios de los grupos son iguales", bold_prefix="H0: ")
add_p(doc, "x̄i = x̄y")
add_p(doc, "Por lo menos los Promedios de 2 grupos son diferentes", bold_prefix="H1: ")
add_p(doc, "x̄i ≠ x̄y")

add_p(doc, bold_prefix="Paso 2: Determinar el nivel de significancia")
add_p(doc, "N.S = 5% (0.05)")

add_p(doc, bold_prefix="Paso 3: Hallar estadísticos descriptivos")
if os.path.exists("generated_kruskal_images/ex3_informe.png"):
    add_img(doc, "generated_kruskal_images/ex3_informe.png", width=Inches(3.1))

add_p(doc, bold_prefix="Paso 4: Prueba de Distribución normal")
add_p(doc, "Los datos de los grupos Aplicación 1, 2 y 3 presentan distribución normal", bold_prefix="H0: ")
add_p(doc, "Por lo menos los datos de un grupo Aplicación 1, 2 o 3 no presentan distribución normal", bold_prefix="H1: ")
if os.path.exists("generated_kruskal_images/ex3_normalidad.png"):
    add_img(doc, "generated_kruskal_images/ex3_normalidad.png", width=Inches(3.5))

add_p(doc, "Dado que los p-valores de los grupos Aplicación 1 (0.189) y Aplicación 2 (0.972) son mayores que el nivel de significancia, pero el grupo Aplicación 3 (0.025) es menor al nivel de significancia por lo que es probable que se rechace la hipótesis nula")
add_p(doc, "A un margen de error de 5 % se concluye que , los datos de los grupos Aplicación 1, 2 y 3, no presentan Distribución normal", bold_prefix="Conclusión: ")

add_p(doc, bold_prefix="Paso 5: Elegir la prueba estadística")
add_p(doc, "Dado que los datos del grupo Aplicación 3 no presentan distribución normal  se aplicara la prueba Kruskal-Wallis (Prueba no paramétrica)")
if os.path.exists("generated_kruskal_images/ex3_rangos.png"):
    add_img(doc, "generated_kruskal_images/ex3_rangos.png", width=Inches(2.8))
if os.path.exists("generated_kruskal_images/ex3_test.png"):
    add_img(doc, "generated_kruskal_images/ex3_test.png", width=Inches(2.4))

add_p(doc, bold_prefix="Paso 6: Contrastación de Hipótesis")
add_p(doc, "Dado que el p -valor de Kruskal -Wallis (0.005) es menor que el nivel de significancia (0.05), se rechaza la hipótesis nula.")

add_p(doc, bold_prefix="Paso 7: Conclusión")
add_p(doc, "A un margen de error del 5% se concluye que, por lo menos los consumos de memoria de 2 aplicaciones son diferentes.", space_after=12)

# ==========================================
# EJERCICIO 4 (RESUELTO)
# ==========================================
add_p(doc, "4. Se desea comparar la latencia (en ms) de tres tipos de redes (X, Y y Z) bajo la misma carga de trabajo. Se han realizado pruebas en diferentes números de redes para cada tipo:\nRed X: 30  32  28  35  33\nRed Y: 45  50  48  47  49  46\nRed Z: 20  22  38  21  25\n¿Existe una diferencia significativa en la latencia de las redes? Contrastar la hipótesis a un margen de error del 4%.", space_after=4)

add_p(doc, bold_prefix="Solución:")
add_p(doc, bold_prefix="Paso 1: Formación de hipótesis")
add_p(doc, "Las Promedios de los grupos son iguales", bold_prefix="H0: ")
add_p(doc, "x̄i = x̄y")
add_p(doc, "Por lo menos los Promedios de 2 grupos son diferentes", bold_prefix="H1: ")
add_p(doc, "x̄i ≠ x̄y")

add_p(doc, bold_prefix="Paso 2: Determinar el nivel de significancia")
add_p(doc, "N.S = 4% (0.04)")

add_p(doc, bold_prefix="Paso 3: Hallar estadísticos descriptivos")
if os.path.exists("generated_kruskal_images/ex4_informe.png"):
    add_img(doc, "generated_kruskal_images/ex4_informe.png", width=Inches(3.1))

add_p(doc, bold_prefix="Paso 4: Prueba de Distribución normal")
add_p(doc, "Los datos de los grupos Red X, Red Y y Red Z presentan distribución normal", bold_prefix="H0: ")
add_p(doc, "Por lo menos los datos de un grupo Red X, Red Y o Red Z no presentan distribución normal", bold_prefix="H1: ")
if os.path.exists("generated_kruskal_images/ex4_normalidad.png"):
    add_img(doc, "generated_kruskal_images/ex4_normalidad.png", width=Inches(3.5))

add_p(doc, "Dado que los p-valores de los grupos Red X (0.980) y Red Y (0.961) son mayores que el nivel de significancia, pero el grupo Red Z (0.040) es menor o igual al nivel de significancia por lo que es probable que se rechace la hipótesis nula")
add_p(doc, "A un margen de error de 4 % se concluye que , los datos de los grupos Red X, Red Y y Red Z, no presentan Distribución normal", bold_prefix="Conclusión: ")

add_p(doc, bold_prefix="Paso 5: Elegir la prueba estadística")
add_p(doc, "Dado que los datos del grupo Red Z no presentan distribución normal  se aplicara la prueba Kruskal-Wallis (Prueba no paramétrica)")
if os.path.exists("generated_kruskal_images/ex4_rangos.png"):
    add_img(doc, "generated_kruskal_images/ex4_rangos.png", width=Inches(2.8))
if os.path.exists("generated_kruskal_images/ex4_test.png"):
    add_img(doc, "generated_kruskal_images/ex4_test.png", width=Inches(2.4))

add_p(doc, bold_prefix="Paso 6: Contrastación de Hipótesis")
add_p(doc, "Dado que el p -valor de Kruskal -Wallis (0.003) es menor que el nivel de significancia (0.04), se rechaza la hipótesis nula.")

add_p(doc, bold_prefix="Paso 7: Conclusión")
add_p(doc, "A un margen de error del 4% se concluye que, por lo menos las latencias de 2 tipos de redes son diferentes.", space_after=12)

# Save DOCX
out_docx = "Ejercicios Kruskal-Wallis_Resuelto_GianFrancoArocutipa.docx"
doc.save(out_docx)
print("Saved DOCX:", out_docx)
