import sys
sys.stdout.reconfigure(encoding='utf-8')
import docx

doc = docx.Document("Ejercicios Pruebas de chi cuadrado.docx")
body = doc._body._element

tbl_counter = 0
p_counter = 0

for i, child in enumerate(body):
    tag = child.tag.split("}")[-1]
    if tag == "tbl":
        t = doc.tables[tbl_counter]
        print(f"[{i}] TABLE {tbl_counter}: {len(t.rows)} rows x {len(t.columns)} cols. Row 0: {[c.text.strip() for c in t.rows[0].cells[:5]]}")
        tbl_counter += 1
    elif tag == "p":
        p = doc.paragraphs[p_counter]
        p_counter += 1
        imgs = len(p._p.xpath('.//@r:embed'))
        t_text = p.text.strip()
        if t_text or imgs:
            print(f"[{i}] P{p_counter-1} (imgs={imgs}): {t_text[:80]}")
        else:
            print(f"[{i}] P{p_counter-1} (empty)")
