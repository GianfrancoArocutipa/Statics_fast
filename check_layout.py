import docx
import os
from PIL import Image

doc = docx.Document('PEREZ_AROCUTIPA_LLICA_APLICACION_ESTADISTICA_INFERENCIAL.docx')
print('File size:', os.path.getsize('PEREZ_AROCUTIPA_LLICA_APLICACION_ESTADISTICA_INFERENCIAL.docx'), 'bytes')
print('Paragraphs:', len(doc.paragraphs))
print('Tables:', len(doc.tables))

drawings = 0
for p in doc.paragraphs:
    drawings += len(p._p.xpath('.//@r:embed'))
for t in doc.tables:
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                drawings += len(p._p.xpath('.//@r:embed'))
print('Total images embedded:', drawings)

# Estimate total vertical height in points
# On Letter page (11 inches = 792 pt, 1 in margins -> 648 pt printable height per page):
# Let's estimate:
total_pt = 0
# Cover page = 1 full page (ends with page break)
# Index page = 1 full page (ends with page break)
# Body pages:
page_breaks = [i for i, p in enumerate(doc.paragraphs) if any('w:br' in r._r.xml and 'page' in r._r.xml for r in p.runs)]
print('Page breaks count:', len(page_breaks))

# Let's calculate the text and table height
body_pt = 0
for p in doc.paragraphs[page_breaks[-1]+1:] if page_breaks else doc.paragraphs:
    # check if image in paragraph
    embeds = p._p.xpath('.//@r:embed')
    if embeds:
        body_pt += 180 # average image height ~2.5 inches = 180 pt
    else:
        text_len = len(p.text)
        if text_len > 0:
            lines = max(1, text_len // 95 + 1)
            body_pt += lines * 14 + p.paragraph_format.space_before.pt + p.paragraph_format.space_after.pt

for t in doc.tables:
    rows_count = len(t.rows)
    body_pt += rows_count * 20 + 15

print('Estimated body height in points:', body_pt)
print('Estimated body pages (body_pt / 648):', body_pt / 648)
print('Total estimated pages (Cover + Index + Body):', 2 + (body_pt / 648))
