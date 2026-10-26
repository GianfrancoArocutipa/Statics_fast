import os
import shutil
import docx
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

original_file = "Ejercicios Pruebas de chi cuadrado.docx"
backup_file = "Ejercicios Pruebas de chi cuadrado_Original.docx"
output_file = "Ejercicios Pruebas de chi cuadrado_Resuelto.docx"

# Backup original
if not os.path.exists(backup_file):
    shutil.copy(original_file, backup_file)
    print(f"Created backup: {backup_file}")

doc = docx.Document(original_file)

def add_p_before(target_p, text, is_empty=False):
    p = target_p.insert_paragraph_before(text if not is_empty else "", style="List Paragraph")
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    for r in p.runs:
        r.font.name = "Calibri"
        r.font.size = Pt(11)
    return p

def add_img_before(target_p, img_path, width):
    p = target_p.insert_paragraph_before("", style="List Paragraph")
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run()
    r.add_picture(img_path, width=width)
    return p

def append_p(doc, text, is_empty=False):
    p = doc.add_paragraph(text if not is_empty else "", style="List Paragraph")
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    for r in p.runs:
        r.font.name = "Calibri"
        r.font.size = Pt(11)
    return p

def append_img(doc, img_path, width):
    p = doc.add_paragraph("", style="List Paragraph")
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run()
    r.add_picture(img_path, width=width)
    return p

# 1. Locate P5 (Enunciado of Exercise 2) to insert Exercise 1 resolution BEFORE it
target_ex2 = None
for p in doc.paragraphs:
    if "En una empresa de desarrollo de software" in p.text:
        target_ex2 = p
        break

if not target_ex2:
    raise Exception("Could not find Exercise 2 paragraph!")

print("Inserting Exercise 1 resolution...")
add_p_before(target_ex2, "Planteamiento de Hipótesis")
add_p_before(target_ex2, "H0: Las proporciones observadas de preferencia por los sistemas operativos en la muestra son iguales a las proporciones de la población (no hay diferencia).")
add_p_before(target_ex2, "p(Windows) = 40%, p(Linux) = 35%, p(macOS) = 25%")
add_p_before(target_ex2, "", is_empty=True)
add_p_before(target_ex2, "H1: Las proporciones observadas de preferencia por los sistemas operativos en la muestra son diferentes de las proporciones de la población (al menos una de las proporciones difiere).")
add_p_before(target_ex2, "", is_empty=True)
add_p_before(target_ex2, "Determinar el nivel de Significancia:5% (0.05)")
add_p_before(target_ex2, "Tabla de distribución de frecuencias")
add_img_before(target_ex2, "generated_spss_images/spss_chi_freq_so.png", width=Inches(3.38))
add_p_before(target_ex2, "", is_empty=True)
add_p_before(target_ex2, "Seleccionar el estadístico de prueba: Prueba no paramétrica Chi-cuadrado de bondad de ajuste")
add_img_before(target_ex2, "generated_spss_images/spss_chi_test_so.png", width=Inches(5.0))
add_p_before(target_ex2, "Contrastación de Hipótesis")
add_p_before(target_ex2, "Dado que p-valor (0.656) es mayor que el nivel de significancia (0.05), es posible que no se rechace la H0")
add_p_before(target_ex2, "Conclusión")
add_p_before(target_ex2, "A un margen del 5% se concluye que, Las proporciones observadas en la muestra no son diferentes de las de la población respecto a la preferencia de los sistemas operativos Windows, Linux y macOS.")
add_p_before(target_ex2, "", is_empty=True)

# 2. Locate P6 ("Prueba binomial") to insert Exercise 2 resolution BEFORE it
target_binomial = None
for p in doc.paragraphs:
    if p.text.strip() == "Prueba binomial":
        target_binomial = p
        break

if not target_binomial:
    raise Exception("Could not find 'Prueba binomial' heading!")

print("Inserting Exercise 2 resolution...")
add_p_before(target_binomial, "Planteamiento de Hipótesis")
add_p_before(target_binomial, "H0: Las proporciones observadas de preferencia por los lenguajes de programación en la muestra son iguales a las proporciones de la población (no hay diferencia).")
add_p_before(target_binomial, "p(Python) = 50%, p(Java) = 35%, p(JavaScript) = 15%")
add_p_before(target_binomial, "", is_empty=True)
add_p_before(target_binomial, "H1: Las proporciones observadas de preferencia por los lenguajes de programación en la muestra son diferentes de las proporciones de la población (al menos una de las proporciones difiere).")
add_p_before(target_binomial, "", is_empty=True)
add_p_before(target_binomial, "Determinar el nivel de Significancia:5% (0.05)")
add_p_before(target_binomial, "Tabla de distribución de frecuencias")
add_img_before(target_binomial, "generated_spss_images/spss_chi_freq_lenguajes.png", width=Inches(3.38))
add_p_before(target_binomial, "", is_empty=True)
add_p_before(target_binomial, "Seleccionar el estadístico de prueba: Prueba no paramétrica Chi-cuadrado de bondad de ajuste")
add_img_before(target_binomial, "generated_spss_images/spss_chi_test_lenguajes.png", width=Inches(5.0))
add_p_before(target_binomial, "Contrastación de Hipótesis")
add_p_before(target_binomial, "Dado que p-valor (1.000) es mayor que el nivel de significancia (0.05), es posible que no se rechace la H0")
add_p_before(target_binomial, "Conclusión")
add_p_before(target_binomial, "A un margen del 5% se concluye que, Las proporciones observadas en la muestra no son diferentes de las de la población respecto a la preferencia de los lenguajes de programación Python, Java y JavaScript.")
add_p_before(target_binomial, "", is_empty=True)

# 3. Correct copy-paste typo in Exercise 4 (Intrusion Detection Algorithm)
print("Correcting Exercise 4 text...")
for p in doc.paragraphs:
    if "p≤75%" in p.text:
        # Check previous paragraph for H0
        pass
    if "El nuevo protocolo no ha aumentado significativamente la tasa de éxito" in p.text:
        # Check if this belongs to intrusion exercise (it has 75% nearby)
        # We can identify by checking surrounding paragraphs or indices
        pass

# Let's find paragraphs belonging to Exercise 4 specifically
for i, p in enumerate(doc.paragraphs):
    if "Un equipo de ingeniería de sistemas está probando un nuevo algoritmo de detección de intrusiones" in p.text:
        # Next paragraphs are H0, H1, etc.
        for j in range(i+1, i+20):
            pj = doc.paragraphs[j]
            if "Prueba de rachas" in pj.text:
                break
            if "El nuevo protocolo no ha aumentado" in pj.text:
                pj.text = pj.text.replace("El nuevo protocolo no ha aumentado significativamente la tasa de éxito en la entrega de paquetes de datos dentro de la red.",
                                          "El nuevo algoritmo no ha mejorado significativamente la tasa de detección de intrusos en la red.")
                for r in pj.runs:
                    r.font.name = "Calibri"
                    r.font.size = Pt(11)
            elif "El nuevo protocolo ha aumentado" in pj.text:
                pj.text = pj.text.replace("El nuevo protocolo ha aumentado significativamente la tasa de éxito en la entrega de paquetes de datos dentro de la red.",
                                          "El nuevo algoritmo ha mejorado significativamente la tasa de detección de intrusos en la red.")
                for r in pj.runs:
                    r.font.name = "Calibri"
                    r.font.size = Pt(11)

# 4. Insert Exercise 6 resolution after Table 1 (at the end of the document)
print("Inserting Exercise 6 resolution...")
append_p(doc, "Planteamiento de Hipótesis")
append_p(doc, "H0: No Existe un patrón o tendencia en el número de correos no deseados que exceden el umbral a lo largo del periodo de 30 días (aleatorio)")
append_p(doc, "", is_empty=True)
append_p(doc, "H1: Existe un patrón o tendencia en el número de correos no deseados que exceden el umbral a lo largo del periodo de 30 días (no aleatorio)")
append_p(doc, "", is_empty=True)
append_p(doc, "Determinar el nivel de Significancia:5% (0.05)")
append_p(doc, "", is_empty=True)
append_p(doc, "Seleccionar el estadístico de prueba: Prueba no paramétrica Rachas")
append_img(doc, "generated_spss_images/spss_rachas_correo.png", width=Inches(2.78))
append_p(doc, "Contrastación de Hipótesis")
append_p(doc, "Dado que p-valor (0.152) es mayor que el nivel de significancia (0.05), es probable que no se rechace la H0")
append_p(doc, "Conclusión")
append_p(doc, "A un margen del 5% se concluye que, No Existe un patrón o tendencia en el número de correos no deseados que exceden el umbral a lo largo del periodo de 30 días (aleatorio)")

# Save completed document
doc.save(output_file)
print(f"Saved completed report to {output_file}")

# Also update the original file directly so the user has it ready
doc.save(original_file)
print(f"Updated original file {original_file}")
