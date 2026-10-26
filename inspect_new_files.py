# -*- coding: utf-8 -*-
import docx
import os
import zipfile

def inspect_file(filepath, out_txt):
    with open(out_txt, 'w', encoding='utf-8') as out:
        if not os.path.exists(filepath):
            out.write(f"File {filepath} not found\n")
            return
        doc = docx.Document(filepath)
        out.write(f"=== {filepath} ===\n")
        out.write(f"Paragraphs: {len(doc.paragraphs)}, Tables: {len(doc.tables)}\n\n")
        for i, p in enumerate(doc.paragraphs):
            drawings = len(p._p.xpath('.//@r:embed'))
            if p.text.strip() or drawings > 0:
                out.write(f"[{i}] {p.text} (drawings: {drawings})\n")
        out.write("\n=== TABLES ===\n")
        for t_idx, table in enumerate(doc.tables):
            out.write(f"\n--- Table {t_idx} ({len(table.rows)} rows x {len(table.columns)} cols) ---\n")
            for r_idx, row in enumerate(table.rows[:10]):
                cells = [c.text.strip().replace('\n', ' ') for c in row.cells]
                out.write(f"  R{r_idx}: {cells}\n")

# Extract media from Reporte_Estadistico (1).docx
report_path = 'Reporte_Estadistico (1).docx'
media_dir = 'extracted_reporte_media'
os.makedirs(media_dir, exist_ok=True)
if os.path.exists(report_path):
    with zipfile.ZipFile(report_path, 'r') as z:
        for f in z.namelist():
            if f.startswith('word/media/'):
                with open(os.path.join(media_dir, os.path.basename(f)), 'wb') as out_f:
                    out_f.write(z.read(f))
                print(f"Extracted media: {f} ({z.getinfo(f).file_size} bytes)")

inspect_file('Ejerc_prueba ANOVA.docx', 'inspect_anova_exercise.txt')
inspect_file('Reporte_Estadistico (1).docx', 'inspect_reporte_estadistico.txt')
print("Inspection complete.")
