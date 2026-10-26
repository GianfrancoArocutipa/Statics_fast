import docx

doc = docx.Document('PEREZ _AROCUTIPA_LLICA - APLICACIÓN DE ESTADÍSTICA INFERENCIAL CON DATOS REALES.docx')

print("Styles in doc:")
for s in doc.styles:
    if s.builtin:
        pass
    else:
        # print custom styles
        # print(s.name)
        pass

# Check sections, margins
for s_idx, sec in enumerate(doc.sections):
    print(f"Section {s_idx}: top={sec.top_margin.pt}pt, bottom={sec.bottom_margin.pt}pt, left={sec.left_margin.pt}pt, right={sec.right_margin.pt}pt, page_w={sec.page_width.pt}pt, page_h={sec.page_height.pt}pt")

# Inspect paragraphs styling, fonts
for i in range(min(40, len(doc.paragraphs))):
    p = doc.paragraphs[i]
    if p.text.strip():
        runs_font = [r.font.name for r in p.runs if r.font.name]
        runs_size = [r.font.size.pt if r.font.size else None for r in p.runs]
        print(f"P{i} [style={p.style.name}, align={p.alignment}]: font={runs_font[:2]}, size={runs_size[:2]} -> {p.text[:60]}")

# Inspect images in model docx
import zipfile
with zipfile.ZipFile('PEREZ _AROCUTIPA_LLICA - APLICACIÓN DE ESTADÍSTICA INFERENCIAL CON DATOS REALES.docx', 'r') as z:
    media_files = [f for f in z.namelist() if f.startswith('word/media/')]
    print(f"Media files in model doc: {len(media_files)}")
    for mf in media_files[:10]:
        print(" ", mf)
